#!/usr/bin/env python3
"""Load only selected P1d posts into memory and remove identity signals."""

from __future__ import annotations

import base64
import binascii
import hashlib
import hmac
import json
import re
import sqlite3
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterator, Mapping, Sequence

from openpyxl import load_workbook


BATCH_SCHEMA_VERSION = "p1d-facilitation-review-v1"
SOURCE_SCHEMA_VERSION = "p1c-role-function-staging-v1"
HMAC_ENV = "CMOOC_RESEARCH_HMAC_KEY_B64"
MAX_CELL_LENGTH = 32_767
REQUIRED_LOGICAL_FILES = {
    "courses_2024",
    "course_memberships_2024",
    "posts_2024",
    "replies_2024",
    "courses_2025",
    "course_memberships_2025",
    "posts_2025",
    "replies_2025",
}

PII_PATTERNS: tuple[tuple[re.Pattern[str], str], ...] = (
    (
        re.compile(r"(?<![\w.+-])[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}(?![\w.-])"),
        "[邮箱已隐藏]",
    ),
    (
        re.compile(r"(?<!\d)(?:\+?86[- ]?)?1[3-9]\d{9}(?!\d)"),
        "[电话已隐藏]",
    ),
    (
        re.compile(r"(?<!\d)\d{17}[\dXx](?!\d)"),
        "[证件号已隐藏]",
    ),
    (
        re.compile(r"(?:微信|vx|wechat|qq)\s*[:：]?\s*[A-Za-z0-9_-]{5,}", re.I),
        "[外部账号已隐藏]",
    ),
    (
        re.compile(
            r"(?:api[_-]?key|access[_-]?token|password|密码)\s*[:=：]\s*\S+",
            re.I,
        ),
        "[凭据已隐藏]",
    ),
    (re.compile(r"https?://[^\s<>\"']+", re.I), "[外部链接]"),
    (re.compile(r"(?<!\w)@[A-Za-z0-9_\-]{2,}"), "[外部账号已隐藏]"),
)


class ReviewContentError(RuntimeError):
    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


@dataclass(frozen=True)
class ReviewContent:
    review_item_id: str
    display_order: int
    title: str
    body: str
    suspected_truncation: bool
    redaction_count: int


class _PlainTextParser(HTMLParser):
    BLOCK_TAGS = {
        "address",
        "article",
        "aside",
        "blockquote",
        "br",
        "div",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "li",
        "p",
        "pre",
        "section",
        "table",
        "tr",
    }

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        normalized = tag.lower()
        if normalized in {"script", "style"}:
            self.ignored_depth += 1
        elif not self.ignored_depth and normalized in self.BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        normalized = tag.lower()
        if normalized in {"script", "style"} and self.ignored_depth:
            self.ignored_depth -= 1
        elif not self.ignored_depth and normalized in self.BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.ignored_depth:
            self.parts.append(data)


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _contains(parent: Path, child: Path) -> bool:
    try:
        child.relative_to(parent)
        return True
    except ValueError:
        return False


def _source_id(value: Any) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        text = str(value)
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value).strip()
    return None if text in {"", "0"} else text


def _post_token(key: bytes, course_id: str, module_id: str) -> str:
    payload = "\x1f".join(("post", course_id, module_id)).encode("utf-8")
    return "post_" + hmac.new(key, payload, hashlib.sha256).hexdigest()


def load_hmac_key(environment: Mapping[str, str] | None = None) -> bytes:
    source = environment if environment is not None else __import__("os").environ
    encoded = source.get(HMAC_ENV, "")
    try:
        key = base64.b64decode(encoded, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise ReviewContentError("HMAC_KEY_INVALID") from exc
    if len(key) < 32:
        raise ReviewContentError("HMAC_KEY_REQUIRED")
    return key


def _load_batch(batch_dir: Path) -> dict[str, Any]:
    try:
        manifest = json.loads(
            (batch_dir / "private-manifest.json").read_text(encoding="utf-8")
        )
    except (OSError, json.JSONDecodeError) as exc:
        raise ReviewContentError("REVIEW_BATCH_INVALID") from exc
    if (
        manifest.get("schema_version") != BATCH_SCHEMA_VERSION
        or manifest.get("identity_blind") is not True
        or manifest.get("content_copied") is not False
        or not isinstance(manifest.get("items"), list)
    ):
        raise ReviewContentError("REVIEW_BATCH_INVALID")
    return manifest


def _load_source_contract(
    research_database: Path, batch: Mapping[str, Any], input_root: Path
) -> dict[str, Path]:
    if not research_database.is_file() or not input_root.is_dir():
        raise ReviewContentError("SOURCE_UNAVAILABLE")
    if _file_sha256(research_database) != batch.get("source_database_sha256"):
        raise ReviewContentError("RESEARCH_DATABASE_CHANGED")
    try:
        connection = sqlite3.connect(
            f"file:{research_database.as_posix()}?mode=ro", uri=True
        )
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA query_only = ON")
        manifests = connection.execute(
            "SELECT dataset_fingerprint, schema_version, learner_projection_enabled "
            "FROM dataset_manifest"
        ).fetchall()
        if len(manifests) != 1:
            raise ReviewContentError("DATASET_MANIFEST_INVALID")
        dataset = manifests[0]
        if (
            dataset["dataset_fingerprint"] != batch.get("dataset_fingerprint")
            or dataset["schema_version"] != SOURCE_SCHEMA_VERSION
            or dataset["learner_projection_enabled"] != 0
        ):
            raise ReviewContentError("DATASET_MANIFEST_INVALID")
        rows = connection.execute(
            "SELECT logical_name, relative_path, sha256 FROM source_file"
        ).fetchall()
    except sqlite3.DatabaseError as exc:
        raise ReviewContentError("RESEARCH_DATABASE_INVALID") from exc
    finally:
        if "connection" in locals():
            connection.close()

    logical_names = {str(row["logical_name"]) for row in rows}
    if logical_names != REQUIRED_LOGICAL_FILES:
        raise ReviewContentError("SOURCE_CONTRACT_INVALID")
    resolved: dict[str, Path] = {}
    canonical_root = input_root.resolve()
    for row in rows:
        path = (canonical_root / str(row["relative_path"])).resolve()
        if not _contains(canonical_root, path) or not path.is_file():
            raise ReviewContentError("SOURCE_CONTRACT_INVALID")
        if _file_sha256(path) != row["sha256"]:
            raise ReviewContentError("SOURCE_FILE_CHANGED")
        resolved[str(row["logical_name"])] = path
    return resolved


def _header_positions(header: Sequence[Any], names: set[str]) -> list[int]:
    return [
        index
        for index, value in enumerate(header)
        if str(value).strip() in names
    ]


def _identity_lexicon(membership_paths: Sequence[Path]) -> tuple[str, ...]:
    identities: set[str] = set()
    for path in membership_paths:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            worksheet = workbook[workbook.sheetnames[0]]
            header = next(
                worksheet.iter_rows(min_row=1, max_row=1, values_only=True)
            )
            positions = _header_positions(
                header, {"user_real_name", "user_nick_name"}
            )
            if not positions:
                raise ReviewContentError("IDENTITY_COLUMNS_MISSING")
            for cells in worksheet.iter_rows(min_row=3, values_only=False):
                if any(cell.data_type == "f" for cell in cells):
                    raise ReviewContentError("XLSX_FORMULA_FORBIDDEN")
                for index in positions:
                    if index >= len(cells):
                        continue
                    value = cells[index].value
                    if value is None:
                        continue
                    text = str(value).strip()
                    if len(text) >= 2 and len(text) <= 128:
                        identities.add(text)
        finally:
            workbook.close()
    return tuple(sorted(identities, key=lambda value: (-len(value), value)))


def _plain_text(value: Any) -> str:
    source = "" if value is None else str(value)
    parser = _PlainTextParser()
    try:
        parser.feed(source)
        parser.close()
        text = "".join(parser.parts)
    except Exception:
        text = re.sub(r"<[^>]*>", " ", source)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\t\f\v]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _redact(text: str, identities: Sequence[str]) -> tuple[str, int]:
    redacted = text
    count = 0
    for identity in identities:
        occurrences = redacted.count(identity)
        if occurrences:
            redacted = redacted.replace(identity, "[姓名或网名已隐藏]")
            count += occurrences
    for pattern, replacement in PII_PATTERNS:
        redacted, substitutions = pattern.subn(replacement, redacted)
        count += substitutions
    return redacted, count


def _selected_post_rows(
    path: Path, key: bytes, selected_event_ids: set[str]
) -> Iterator[tuple[str, str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        header = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        indices = {
            str(value).strip(): index for index, value in enumerate(header)
        }
        required = {"module_id", "course_id", "module_title", "module_content"}
        if not required.issubset(indices):
            raise ReviewContentError("POST_COLUMNS_MISSING")
        for cells in worksheet.iter_rows(min_row=3, values_only=False):
            if any(cell.data_type == "f" for cell in cells):
                raise ReviewContentError("XLSX_FORMULA_FORBIDDEN")
            module_id = _source_id(cells[indices["module_id"]].value)
            course_id = _source_id(cells[indices["course_id"]].value)
            if not module_id or not course_id:
                continue
            source_event_id = _post_token(key, course_id, module_id)
            if source_event_id not in selected_event_ids:
                continue
            title = "" if cells[indices["module_title"]].value is None else str(
                cells[indices["module_title"]].value
            )
            body = "" if cells[indices["module_content"]].value is None else str(
                cells[indices["module_content"]].value
            )
            yield source_event_id, title, body
    finally:
        workbook.close()


def load_review_contents(
    *,
    batch_dir_value: str,
    research_database_value: str,
    input_root_value: str,
    key: bytes,
    selected_review_ids: set[str] | None = None,
) -> tuple[dict[str, Any], dict[str, ReviewContent]]:
    batch_dir = Path(batch_dir_value).expanduser().resolve()
    research_database = Path(research_database_value).expanduser().resolve()
    input_root = Path(input_root_value).expanduser().resolve()
    batch = _load_batch(batch_dir)
    source_paths = _load_source_contract(research_database, batch, input_root)
    identities = _identity_lexicon(
        (
            source_paths["course_memberships_2024"],
            source_paths["course_memberships_2025"],
        )
    )

    item_by_event: dict[str, Mapping[str, Any]] = {}
    all_review_ids: set[str] = set()
    all_source_event_ids: set[str] = set()
    for item in batch["items"]:
        source_event_id = str(item.get("source_event_id", ""))
        review_item_id = str(item.get("review_item_id", ""))
        if (
            not source_event_id.startswith("post_")
            or not review_item_id.startswith("review_")
            or review_item_id in all_review_ids
            or source_event_id in all_source_event_ids
        ):
            raise ReviewContentError("REVIEW_BATCH_INVALID")
        all_review_ids.add(review_item_id)
        all_source_event_ids.add(source_event_id)
        if selected_review_ids is None or review_item_id in selected_review_ids:
            if source_event_id in item_by_event:
                raise ReviewContentError("REVIEW_BATCH_INVALID")
            item_by_event[source_event_id] = item
    if selected_review_ids is not None and (
        not selected_review_ids or not selected_review_ids.issubset(all_review_ids)
    ):
        raise ReviewContentError("REVIEW_SELECTION_INVALID")

    contents: dict[str, ReviewContent] = {}
    for path in (source_paths["posts_2024"], source_paths["posts_2025"]):
        for source_event_id, raw_title, raw_body in _selected_post_rows(
            path, key, set(item_by_event)
        ):
            item = item_by_event.get(source_event_id)
            if item is None:
                continue
            title, title_redactions = _redact(_plain_text(raw_title), identities)
            body, body_redactions = _redact(_plain_text(raw_body), identities)
            review_item_id = str(item["review_item_id"])
            contents[review_item_id] = ReviewContent(
                review_item_id=review_item_id,
                display_order=int(item["display_order"]),
                title=title or "（无标题）",
                body=body or "（正文为空）",
                suspected_truncation=max(len(raw_title), len(raw_body))
                >= MAX_CELL_LENGTH,
                redaction_count=title_redactions + body_redactions,
            )
            if len(contents) == len(item_by_event):
                break
        if len(contents) == len(item_by_event):
            break
    if len(contents) != len(item_by_event):
        raise ReviewContentError("SELECTED_CONTENT_NOT_FOUND")

    public_batch = {
        "batch_id": batch["batch_id"],
        "schema_version": batch["schema_version"],
        "selected_count": len(contents),
        "population_count": len(batch["items"]),
        "selection_applied": selected_review_ids is not None,
        "codebook_version": batch["codebook_version"],
        "identity_blind": True,
        "content_copied": False,
    }
    return public_batch, contents
