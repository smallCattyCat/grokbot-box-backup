#!/usr/bin/env python3
"""Build and evaluate leakage-safe cMOOC intervention decision snapshots.

The raw workbooks are read-only inputs. Formal outputs contain only HMAC-backed
stable keys, timestamps, numeric features, and controlled labels. No forum text,
raw platform identifiers, membership attributes, or final-like counts are
written to the formal artifact set.
"""

from __future__ import annotations

import argparse
import base64
import bisect
import csv
import hashlib
import hmac
import json
import math
import os
import secrets
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from decision_reference_integrity import (
    REFERENCE_BLIND_ORDER_VERSION,
    REFERENCE_INTEGRITY_VERSION,
    ReferenceIntegrityError,
    has_complete_reference_integrity,
    reference_blind_order_key,
    reference_integrity,
    verify_reference_integrity,
)
from research_decision.benchmark import evaluate_replay_benchmark
from research_decision.oracle import (
    ORACLE_RESULT_FIELDS,
    compare_point_in_time_implementations,
    sql_point_in_time_recompute,
)
from research_decision.temporal_contract import (
    audit_temporal_rows,
    contract_manifest,
)


SCHEMA_VERSION = "cmooc-decision-snapshot-v1"
BENCHMARK_SCHEMA_VERSION = "cmooc-scheduler-benchmark-v2"
REFERENCE_SCHEMA_VERSION = "cmooc-decision-reference-selection-v1"
DEVELOPMENT_YEAR = 2024
HOLDOUT_YEAR = 2025
DECISION_QUANTILE = 0.75
OUTCOME_QUANTILE = 0.90
DEFAULT_BUDGETS = (1, 3, 5)
DEFAULT_REFERENCE_LIMIT = 60
SOURCE_TIME_ASSUMPTION = "naive_platform_local_time_no_dst_adjustment"
HISTORY_FEATURE_AUDIT_FIELDS = (
    "author_prior_post_count_asof",
    "author_prior_external_response_count_asof",
    "author_prior_external_response_rate_asof",
    "author_prior_matured_gap_count_asof",
    "course_post_count_prior_7d_asof",
    "course_external_comment_count_prior_7d_asof",
    "source_observation_end_after_decision_microseconds",
)

FORMAL_SNAPSHOT_FIELDS = (
    "schema_version",
    "snapshot_key",
    "thread_key",
    "course_key",
    "author_key",
    "year",
    "decision_at",
    "decision_date",
    "decision_lag_hours",
    "outcome_window_end_at",
    "outcome_window_hours",
    "comment_count_asof",
    "external_comment_count_asof",
    "self_comment_count_asof",
    "unique_external_author_count_asof",
    "reply_edge_count_asof",
    "max_reply_depth_asof",
    "future_parent_reference_count_asof",
    "author_prior_post_count_asof",
    "author_prior_external_response_count_asof",
    "author_prior_external_response_rate_asof",
    "author_prior_matured_gap_count_asof",
    "course_post_count_prior_7d_asof",
    "course_external_comment_count_prior_7d_asof",
    "decision_eligibility",
    "suggested_function",
    "outcome_observable",
    "source_observation_end_after_decision_microseconds",
    "first_external_response_after_decision_microseconds",
    "organic_external_response_in_window",
    "external_response_count_in_window",
    "unique_external_responder_count_in_window",
    "support_gap_at_horizon",
    "feature_hash",
    "asof_boundary_ok",
)

REFERENCE_LABEL_FIELDS = ("intervene", "main_function", "target_object")
REFERENCE_FIELDS = (
    "reference_id",
    "snapshot_key",
    "thread_key",
    "course_key",
    "author_key",
    "year",
    "decision_at",
    "sampling_stratum",
    "sampling_reason",
    *REFERENCE_LABEL_FIELDS,
)

BENCHMARK_ACTION_FIELDS = (
    "year",
    "split",
    "policy",
    "daily_budget_k",
    "snapshot_key",
    "thread_key",
    "course_key",
    "author_key",
    "decision_at",
    "tick_at",
    "selection_delay_minutes",
    "selection_rank_at_tick",
)

FORBIDDEN_FORMAL_TOKENS = (
    "content",
    "title",
    "module_id",
    "comment_id",
    "course_id",
    "author_id",
    "real_name",
    "nickname",
    "like_count",
)


class DecisionSnapshotError(RuntimeError):
    pass


@dataclass(frozen=True)
class RawPost:
    year: int
    course_id: str
    module_id: str
    author_id: str
    created_at: datetime

    @property
    def raw_key(self) -> tuple[int, str, str]:
        return (self.year, self.course_id, self.module_id)


@dataclass(frozen=True)
class RawComment:
    year: int
    course_id: str
    module_id: str
    comment_id: str
    author_id: str
    created_at: datetime
    is_reply: bool
    parent_comment_id: str

    @property
    def post_key(self) -> tuple[int, str, str]:
        return (self.year, self.course_id, self.module_id)


@dataclass(frozen=True)
class TimePolicy:
    development_year: int
    decision_quantile: float
    outcome_quantile: float
    raw_decision_hours: float
    raw_outcome_end_hours: float
    decision_lag_hours: int
    outcome_end_hours: int
    calibration_count: int

    @property
    def outcome_window_hours(self) -> int:
        return self.outcome_end_hours - self.decision_lag_hours


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[1]


def _default_input_root() -> Path:
    configured = os.environ.get("CMOOC_RESEARCH_INPUT_ROOT")
    if configured:
        return Path(configured).expanduser().resolve()
    return _repo_root().parent.parent / "cmooc研究数据24-25"


def _research_home() -> Path:
    configured = os.environ.get("CMOOC_RESEARCH_HOME")
    if configured:
        return Path(configured).expanduser().resolve()
    local_app_data = os.environ.get("LOCALAPPDATA")
    if not local_app_data:
        raise DecisionSnapshotError("LOCAL_RESEARCH_HOME_REQUIRED")
    return Path(local_app_data).resolve() / "cmooc-research"


def _default_output_dir() -> Path:
    return _research_home() / "p3a" / "decision-snapshot-v1"


def _contains(parent: Path, child: Path) -> bool:
    try:
        child.relative_to(parent)
        return True
    except ValueError:
        return False


def _safe_output_dir(value: str | None) -> Path:
    root = _research_home()
    output = Path(value).expanduser().resolve() if value else _default_output_dir()
    if (
        output == root
        or not _contains(root, output)
        or _contains(_repo_root(), output)
        or str(output).startswith("\\\\")
    ):
        raise DecisionSnapshotError("OUTPUT_OUTSIDE_LOCAL_RESEARCH_HOME")
    return output


def _normalize_id(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_timestamp(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    text = _normalize_id(value)
    if not text:
        return None
    text = text.replace("/", "-").replace("T", " ").strip()
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def _iso(value: datetime) -> str:
    return value.isoformat(timespec="seconds")


def _timedelta_microseconds(value: timedelta) -> int:
    return (
        (value.days * 86_400 + value.seconds) * 1_000_000
        + value.microseconds
    )


def _quantile(values: Sequence[float], probability: float) -> float:
    if not values:
        raise DecisionSnapshotError("NO_DEVELOPMENT_FIRST_RESPONSES")
    ordered = sorted(values)
    position = (len(ordered) - 1) * probability
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return ordered[lower]
    return (
        ordered[lower] * (upper - position)
        + ordered[upper] * (position - lower)
    )


def _round_nearest_hour(value: float) -> int:
    return max(1, int(math.floor(value + 0.5)))


def _stable_key(secret: bytes, namespace: str, *parts: object) -> str:
    message = "\x1f".join([namespace, *(str(part) for part in parts)])
    return hmac.new(secret, message.encode("utf-8"), hashlib.sha256).hexdigest()[:24]


def _canonical_hash(value: Mapping[str, Any]) -> str:
    encoded = json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _source_files(input_root: Path) -> dict[int, dict[str, Path]]:
    raw = input_root / "01_原始数据" / "数据统计"
    result = {
        year: {
            "posts": raw / f"{year}图文.xlsx",
            "comments": raw / f"{year}图文评论.xlsx",
            "courses": raw / f"{year}课程.xlsx",
            "members": raw / f"{year}课程成员.xlsx",
        }
        for year in (DEVELOPMENT_YEAR, HOLDOUT_YEAR)
    }
    missing = [
        str(path)
        for files in result.values()
        for path in files.values()
        if not path.is_file()
    ]
    if missing:
        raise DecisionSnapshotError("SOURCE_FILES_MISSING:" + "|".join(missing))
    return result


def _iter_sheet(path: Path) -> Iterable[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [
        str(value).strip() if value is not None else "" for value in next(rows)
    ]
    for values in rows:
        yield dict(zip(headers, values))
    workbook.close()


def load_raw_events(
    input_root: Path,
) -> tuple[list[RawPost], list[RawComment], dict[str, Any]]:
    files = _source_files(input_root)
    posts: list[RawPost] = []
    comments: list[RawComment] = []
    post_keys: set[tuple[int, str, str]] = set()
    comment_keys: set[tuple[int, str, str, str]] = set()
    raw_stats: dict[str, Any] = {
        "schema_version": SCHEMA_VERSION,
        "source_time_assumption": SOURCE_TIME_ASSUMPTION,
        "years": {},
        "excluded_final_snapshot_fields": [
            "comment_like_count",
            "final_comment_count",
            "final_reply_count",
            "final_network_centrality",
            "post_or_comment_text",
            "membership_role_as_primary_feature",
        ],
    }

    for year, year_files in files.items():
        year_stats: Counter[str] = Counter()
        for row in _iter_sheet(year_files["posts"]):
            year_stats["post_rows"] += 1
            course_id = _normalize_id(row.get("course_id"))
            module_id = _normalize_id(row.get("module_id"))
            author_id = _normalize_id(row.get("created_by"))
            created_at = _parse_timestamp(row.get("created_time"))
            if not course_id or not module_id:
                year_stats["post_missing_stable_key"] += 1
                continue
            key = (year, course_id, module_id)
            if key in post_keys:
                year_stats["duplicate_post_key"] += 1
                continue
            post_keys.add(key)
            if created_at is None:
                year_stats["post_time_unparseable"] += 1
                continue
            if not author_id:
                year_stats["post_author_missing"] += 1
            posts.append(
                RawPost(
                    year=year,
                    course_id=course_id,
                    module_id=module_id,
                    author_id=author_id,
                    created_at=created_at,
                )
            )

        for row in _iter_sheet(year_files["comments"]):
            year_stats["comment_rows"] += 1
            course_id = _normalize_id(row.get("course_id"))
            module_id = _normalize_id(row.get("module_id"))
            comment_id = _normalize_id(row.get("comment_id"))
            author_id = _normalize_id(row.get("created_by"))
            created_at = _parse_timestamp(row.get("created_time"))
            key = (year, course_id, module_id, comment_id)
            if not all((course_id, module_id, comment_id)):
                year_stats["comment_missing_stable_key"] += 1
                continue
            if key in comment_keys:
                year_stats["duplicate_comment_key"] += 1
                continue
            comment_keys.add(key)
            if created_at is None:
                year_stats["comment_time_unparseable"] += 1
                continue
            if not author_id:
                year_stats["comment_author_missing"] += 1
            is_reply = _normalize_id(row.get("comment_reply")) == "1"
            parent = _normalize_id(row.get("comment_reply_child_id"))
            if parent in {"", "-1"}:
                parent = _normalize_id(row.get("comment_reply_id"))
            comments.append(
                RawComment(
                    year=year,
                    course_id=course_id,
                    module_id=module_id,
                    comment_id=comment_id,
                    author_id=author_id,
                    created_at=created_at,
                    is_reply=is_reply,
                    parent_comment_id=parent if parent not in {"-1"} else "",
                )
            )

        # Membership is audited only as a snapshot sensitivity source.
        member_rows = 0
        member_course_user_keys: set[tuple[str, str]] = set()
        member_duplicate_keys = 0
        for row in _iter_sheet(year_files["members"]):
            member_rows += 1
            key = (
                _normalize_id(row.get("course_id")),
                _normalize_id(row.get("user_id")),
            )
            if key in member_course_user_keys:
                member_duplicate_keys += 1
            member_course_user_keys.add(key)
        year_stats["membership_rows"] = member_rows
        year_stats["membership_duplicate_course_user_keys"] = member_duplicate_keys
        raw_stats["years"][str(year)] = dict(year_stats)

    post_by_key = {post.raw_key: post for post in posts}
    comment_ids_by_post: dict[tuple[int, str, str], set[str]] = defaultdict(set)
    for item in comments:
        comment_ids_by_post[item.post_key].add(item.comment_id)

    relation_stats: Counter[str] = Counter()
    for item in comments:
        target = post_by_key.get(item.post_key)
        if target is None:
            relation_stats["orphan_module_comment"] += 1
            continue
        if item.created_at < target.created_at:
            relation_stats["comment_before_post"] += 1
        if item.is_reply and item.parent_comment_id:
            relation_stats["reply_parent_expected"] += 1
            if item.parent_comment_id in comment_ids_by_post[item.post_key]:
                relation_stats["reply_parent_resolved_same_thread"] += 1
            else:
                relation_stats["reply_parent_unresolved"] += 1
    raw_stats["relationships"] = dict(relation_stats)
    raw_stats["usable_post_count"] = len(posts)
    raw_stats["usable_comment_count"] = len(comments)
    raw_stats["source_files"] = {
        path.name: {
            "bytes": path.stat().st_size,
            "sha256": _sha256_file(path),
        }
        for year_files in files.values()
        for path in year_files.values()
    }
    return posts, comments, raw_stats


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _comments_by_post(
    comments: Iterable[RawComment],
) -> dict[tuple[int, str, str], list[RawComment]]:
    result: dict[tuple[int, str, str], list[RawComment]] = defaultdict(list)
    for item in comments:
        result[item.post_key].append(item)
    for values in result.values():
        values.sort(key=lambda item: (item.created_at, item.comment_id))
    return result


def _external_comments(
    post: RawPost, comments: Iterable[RawComment]
) -> list[RawComment]:
    return [
        item
        for item in comments
        if item.author_id and item.author_id != post.author_id
        and item.created_at >= post.created_at
    ]


def calibrate_time_policy(
    posts: Sequence[RawPost],
    comments: Sequence[RawComment],
    *,
    development_year: int = DEVELOPMENT_YEAR,
) -> TimePolicy:
    by_post = _comments_by_post(comments)
    first_response_hours: list[float] = []
    for post in posts:
        if post.year != development_year:
            continue
        external = _external_comments(post, by_post.get(post.raw_key, ()))
        if external:
            first_response_hours.append(
                (external[0].created_at - post.created_at).total_seconds() / 3600
            )
    raw_decision = _quantile(first_response_hours, DECISION_QUANTILE)
    raw_outcome = _quantile(first_response_hours, OUTCOME_QUANTILE)
    decision_hours = _round_nearest_hour(raw_decision)
    outcome_hours = _round_nearest_hour(raw_outcome)
    if outcome_hours <= decision_hours:
        outcome_hours = decision_hours + 1
    return TimePolicy(
        development_year=development_year,
        decision_quantile=DECISION_QUANTILE,
        outcome_quantile=OUTCOME_QUANTILE,
        raw_decision_hours=raw_decision,
        raw_outcome_end_hours=raw_outcome,
        decision_lag_hours=decision_hours,
        outcome_end_hours=outcome_hours,
        calibration_count=len(first_response_hours),
    )


def _course_observation_ends(
    posts: Sequence[RawPost], comments: Sequence[RawComment]
) -> dict[tuple[int, str], datetime]:
    result: dict[tuple[int, str], datetime] = {}
    for item in [*posts, *comments]:
        key = (item.year, item.course_id)
        result[key] = max(result.get(key, item.created_at), item.created_at)
    return result


def _reply_depth(
    asof_comments: Sequence[RawComment],
) -> tuple[int, int]:
    asof_by_id = {item.comment_id: item for item in asof_comments}
    memo: dict[str, int] = {}
    visiting: set[str] = set()

    def depth(comment_id: str) -> int:
        if comment_id in memo:
            return memo[comment_id]
        item = asof_by_id.get(comment_id)
        if item is None or not item.parent_comment_id:
            return 0
        if comment_id in visiting:
            return 0
        parent = asof_by_id.get(item.parent_comment_id)
        if parent is None:
            return 0
        visiting.add(comment_id)
        value = depth(parent.comment_id) + 1
        visiting.discard(comment_id)
        memo[comment_id] = value
        return value

    maximum = max((depth(item.comment_id) for item in asof_comments), default=0)
    # The legacy field name is retained for artifact compatibility. Its value
    # is strictly the number of parent references unresolved at decision_at;
    # it never checks whether a parent appears in a later export event.
    unresolved_parent_count = sum(
        1
        for item in asof_comments
        if item.parent_comment_id
        and item.parent_comment_id not in asof_by_id
    )
    return maximum, unresolved_parent_count


def build_snapshot_records(
    posts: Sequence[RawPost],
    comments: Sequence[RawComment],
    policy: TimePolicy,
    secret: bytes,
    *,
    observation_ends: Mapping[tuple[int, str], datetime] | None = None,
) -> list[dict[str, Any]]:
    by_post = _comments_by_post(comments)
    post_by_key = {item.raw_key: item for item in posts}
    valid_comments = [
        item
        for item in comments
        if item.post_key in post_by_key
        and item.created_at >= post_by_key[item.post_key].created_at
    ]
    by_post = _comments_by_post(valid_comments)

    author_posts: dict[str, list[RawPost]] = defaultdict(list)
    course_post_times: dict[tuple[int, str], list[datetime]] = defaultdict(list)
    course_external_times: dict[tuple[int, str], list[datetime]] = defaultdict(list)
    first_external_by_post: dict[tuple[int, str, str], datetime | None] = {}
    for post in posts:
        author_posts[post.author_id].append(post)
        course_post_times[(post.year, post.course_id)].append(post.created_at)
        external = _external_comments(post, by_post.get(post.raw_key, ()))
        first_external_by_post[post.raw_key] = (
            external[0].created_at if external else None
        )
        course_external_times[(post.year, post.course_id)].extend(
            item.created_at for item in external
        )
    for values in author_posts.values():
        values.sort(key=lambda item: item.created_at)
    for mapping in (course_post_times, course_external_times):
        for values in mapping.values():
            values.sort()

    records: list[dict[str, Any]] = []
    for post in sorted(posts, key=lambda item: (item.created_at, item.raw_key)):
        decision_at = post.created_at + timedelta(hours=policy.decision_lag_hours)
        outcome_end = post.created_at + timedelta(hours=policy.outcome_end_hours)
        thread_comments = by_post.get(post.raw_key, [])
        asof = [item for item in thread_comments if item.created_at <= decision_at]
        outcome = [
            item
            for item in thread_comments
            if decision_at < item.created_at <= outcome_end
        ]
        asof_external = _external_comments(post, asof)
        outcome_external = _external_comments(post, outcome)
        self_asof = [
            item for item in asof if item.author_id == post.author_id
        ]
        max_depth, unresolved_parent_count = _reply_depth(asof)

        prior_posts = [
            item
            for item in author_posts.get(post.author_id, ())
            if item.raw_key != post.raw_key and item.created_at < decision_at
        ]
        prior_response_count = sum(
            1
            for item in prior_posts
            if first_external_by_post.get(item.raw_key) is not None
            and first_external_by_post[item.raw_key] <= decision_at  # type: ignore[operator]
        )
        prior_matured_gap_count = sum(
            1
            for item in prior_posts
            if item.created_at + timedelta(hours=policy.outcome_end_hours)
            <= decision_at
            and (
                first_external_by_post.get(item.raw_key) is None
                or first_external_by_post[item.raw_key]
                > item.created_at
                + timedelta(hours=policy.outcome_end_hours)  # type: ignore[operator]
            )
        )
        prior_response_rate = (
            prior_response_count / len(prior_posts) if prior_posts else 0.0
        )

        course_key_raw = (post.year, post.course_id)
        window_start = decision_at - timedelta(days=7)
        post_times = course_post_times[course_key_raw]
        external_times = course_external_times[course_key_raw]
        post_window_lower = bisect.bisect_right(
            post_times, window_start
        )
        post_window_upper = bisect.bisect_right(
            post_times, decision_at
        )
        external_window_lower = bisect.bisect_right(
            external_times, window_start
        )
        external_window_upper = bisect.bisect_right(
            external_times, decision_at
        )
        course_posts_prior_7d = (
            post_window_upper - post_window_lower
        )
        course_external_prior_7d = (
            external_window_upper - external_window_lower
        )

        source_observation_end = (
            outcome_end
            if observation_ends is None
            else min(
                outcome_end,
                observation_ends.get(
                    course_key_raw, post.created_at
                ),
            )
        )
        observable = int(source_observation_end >= outcome_end)
        source_observation_end_after_decision_microseconds = (
            _timedelta_microseconds(
                source_observation_end - decision_at
            )
        )
        already_responded = bool(asof_external)
        eligibility = (
            "already_responded" if already_responded else "eligible_unanswered"
        )
        organic_response = int(bool(outcome_external)) if observable else 0
        first_external_response_after_decision_microseconds = None
        if outcome_external:
            first_external_response_after_decision_microseconds = (
                _timedelta_microseconds(
                    outcome_external[0].created_at - decision_at
                )
            )
        support_gap = (
            int(not already_responded and not outcome_external)
            if observable
            else 0
        )
        snapshot_key = _stable_key(
            secret,
            "snapshot",
            post.year,
            post.course_id,
            post.module_id,
            _iso(decision_at),
        )
        feature_values = {
            "decision_lag_hours": policy.decision_lag_hours,
            "comment_count_asof": len(asof),
            "external_comment_count_asof": len(asof_external),
            "self_comment_count_asof": len(self_asof),
            "unique_external_author_count_asof": len(
                {item.author_id for item in asof_external}
            ),
            "reply_edge_count_asof": sum(
                1
                for item in asof
                if item.parent_comment_id
                and any(
                    parent.comment_id == item.parent_comment_id
                    for parent in asof
                )
            ),
            "max_reply_depth_asof": max_depth,
            "future_parent_reference_count_asof": unresolved_parent_count,
            "author_prior_post_count_asof": len(prior_posts),
            "author_prior_external_response_count_asof": prior_response_count,
            "author_prior_external_response_rate_asof": round(
                prior_response_rate, 6
            ),
            "author_prior_matured_gap_count_asof": prior_matured_gap_count,
            "course_post_count_prior_7d_asof": course_posts_prior_7d,
            "course_external_comment_count_prior_7d_asof": (
                course_external_prior_7d
            ),
        }
        asof_boundary_ok = int(
            post.created_at <= decision_at
            and all(
                item.created_at <= decision_at for item in asof
            )
            and all(
                decision_at < item.created_at <= outcome_end
                for item in outcome
            )
            and all(
                item.created_at < decision_at for item in prior_posts
            )
            and all(
                window_start < value <= decision_at
                for value in post_times[
                    post_window_lower:post_window_upper
                ]
            )
            and all(
                window_start < value <= decision_at
                for value in external_times[
                    external_window_lower:external_window_upper
                ]
            )
            and (
                first_external_response_after_decision_microseconds
                is None
                or 0
                < first_external_response_after_decision_microseconds
                <= _timedelta_microseconds(
                    outcome_end - decision_at
                )
            )
        )
        record = {
            "schema_version": SCHEMA_VERSION,
            "snapshot_key": snapshot_key,
            "thread_key": _stable_key(
                secret, "thread", post.year, post.course_id, post.module_id
            ),
            "course_key": _stable_key(
                secret, "course", post.year, post.course_id
            ),
            "author_key": _stable_key(secret, "author", post.author_id),
            "year": post.year,
            "decision_at": _iso(decision_at),
            "decision_date": decision_at.date().isoformat(),
            **feature_values,
            "outcome_window_end_at": _iso(outcome_end),
            "outcome_window_hours": policy.outcome_window_hours,
            "decision_eligibility": eligibility,
            # Snapshot-only text cannot safely support a function inference.
            "suggested_function": "undetermined",
            "outcome_observable": observable,
            "source_observation_end_after_decision_microseconds": (
                source_observation_end_after_decision_microseconds
            ),
            "first_external_response_after_decision_microseconds": (
                first_external_response_after_decision_microseconds
            ),
            "organic_external_response_in_window": organic_response,
            "external_response_count_in_window": (
                len(outcome_external) if observable else 0
            ),
            "unique_external_responder_count_in_window": (
                len({item.author_id for item in outcome_external})
                if observable
                else 0
            ),
            "support_gap_at_horizon": support_gap,
            "feature_hash": _canonical_hash(feature_values),
            "asof_boundary_ok": asof_boundary_ok,
        }
        _validate_formal_record(record)
        records.append(record)
    return records


def _validate_formal_record(record: Mapping[str, Any]) -> None:
    if set(record) != set(FORMAL_SNAPSHOT_FIELDS):
        missing = sorted(set(FORMAL_SNAPSHOT_FIELDS) - set(record))
        extra = sorted(set(record) - set(FORMAL_SNAPSHOT_FIELDS))
        raise DecisionSnapshotError(
            f"FORMAL_SCHEMA_MISMATCH:missing={missing}:extra={extra}"
        )
    lowered = [field.lower() for field in record]
    if any(
        token in field
        for field in lowered
        for token in FORBIDDEN_FORMAL_TOKENS
    ):
        raise DecisionSnapshotError("FORBIDDEN_FORMAL_FIELD")
    decision_at = datetime.fromisoformat(str(record["decision_at"]))
    outcome_end = datetime.fromisoformat(str(record["outcome_window_end_at"]))
    if outcome_end <= decision_at:
        raise DecisionSnapshotError("INVALID_OUTCOME_WINDOW")
    if int(record["asof_boundary_ok"]) != 1:
        raise DecisionSnapshotError("ASOF_BOUNDARY_INVARIANT_FAILED")


def _point_in_time_oracle_audit(
    posts: Sequence[RawPost],
    comments: Sequence[RawComment],
    records: Sequence[Mapping[str, Any]],
    policy: TimePolicy,
    secret: bytes,
) -> dict[str, Any]:
    """Compare formal snapshots with an independent local SQL recomputation."""

    record_by_snapshot = {
        str(row["snapshot_key"]): row
        for row in records
        if int(row.get("outcome_observable", 0) or 0) == 1
    }
    post_by_raw_key = {post.raw_key: post for post in posts}
    valid_comments = [
        item
        for item in comments
        if item.post_key in post_by_raw_key
        and item.created_at >= post_by_raw_key[item.post_key].created_at
    ]
    comments_by_post = _comments_by_post(valid_comments)
    oracle_decisions: list[dict[str, Any]] = []
    oracle_events: list[dict[str, Any]] = []

    for post in sorted(posts, key=lambda item: (item.created_at, item.raw_key)):
        decision_at = post.created_at + timedelta(
            hours=policy.decision_lag_hours
        )
        snapshot_key = _stable_key(
            secret,
            "snapshot",
            post.year,
            post.course_id,
            post.module_id,
            _iso(decision_at),
        )
        record = record_by_snapshot.get(snapshot_key)
        if record is None:
            continue
        thread_key = str(record["thread_key"])
        oracle_decisions.append(
            {
                "snapshot_key": snapshot_key,
                "thread_key": thread_key,
                "decision_at": record["decision_at"],
                "outcome_window_end_at": record[
                    "outcome_window_end_at"
                ],
            }
        )
        thread_comments = sorted(
            comments_by_post.get(post.raw_key, ()),
            key=lambda item: (
                item.created_at,
                item.comment_id,
                item.author_id,
                item.parent_comment_id,
            ),
        )
        parent_at = {
            item.comment_id: item.created_at for item in thread_comments
        }
        for index, item in enumerate(thread_comments):
            parent_visible_at_decision = bool(
                item.parent_comment_id
                and item.parent_comment_id in parent_at
                and parent_at[item.parent_comment_id] <= decision_at
            )
            oracle_events.append(
                {
                    "event_key": _stable_key(
                        secret,
                        "oracle-event",
                        snapshot_key,
                        item.comment_id,
                        _iso(item.created_at),
                        index,
                    ),
                    "thread_key": thread_key,
                    "actor_key": _stable_key(
                        secret, "oracle-actor", item.author_id
                    ),
                    "event_time": _iso(item.created_at),
                    "is_external": int(item.author_id != post.author_id),
                    "is_reply": int(parent_visible_at_decision),
                }
            )

    implementation_audit = compare_point_in_time_implementations(
        oracle_decisions,
        oracle_events,
        preferred_backend="auto",
    )
    sql_rows, backend = sql_point_in_time_recompute(
        oracle_decisions,
        oracle_events,
        preferred_backend="auto",
    )
    formal_by_key = {
        str(row["snapshot_key"]): row
        for row in record_by_snapshot.values()
    }
    mismatches_by_field: Counter[str] = Counter()
    for oracle_row in sql_rows:
        formal = formal_by_key[str(oracle_row["snapshot_key"])]
        for field in ORACLE_RESULT_FIELDS[1:]:
            formal_value = formal[field]
            if (
                field
                == "first_external_response_after_decision_microseconds"
                and formal_value in (None, "")
            ):
                formal_value = -1
            if int(formal_value) != int(oracle_row[field]):
                mismatches_by_field[field] += 1
    formal_mismatch_count = sum(mismatches_by_field.values())
    formal_projection = [
        {
            field: (
                str(row[field])
                if field == "snapshot_key"
                else (
                    -1
                    if field
                    == "first_external_response_after_decision_microseconds"
                    and row[field] in (None, "")
                    else int(row[field])
                )
            )
            for field in ORACLE_RESULT_FIELDS
        }
        for row in sorted(
            record_by_snapshot.values(),
            key=lambda item: str(item["snapshot_key"]),
        )
    ]
    result = {
        **implementation_audit,
        "status": (
            "passed"
            if implementation_audit["status"] == "passed"
            and formal_mismatch_count == 0
            else "failed"
        ),
        "contains_row_level_data": False,
        "formal_snapshot_compared_row_count": len(sql_rows),
        "formal_snapshot_mismatch_count": formal_mismatch_count,
        "formal_snapshot_mismatches_by_field": dict(
            sorted(mismatches_by_field.items())
        ),
        "formal_snapshot_result_hash": _canonical_hash(
            {"rows": formal_projection}
        ),
        "backend": backend.to_dict(),
        "scope": (
            "comment/reply structural features and subsequent external "
            "response outcomes; raw events remain in memory"
        ),
    }
    result.pop("audit_hash", None)
    result["audit_hash"] = _canonical_hash(result)
    if result["status"] != "passed":
        raise DecisionSnapshotError("POINT_IN_TIME_ORACLE_MISMATCH")
    return result


def _history_feature_audit(
    posts: Sequence[RawPost],
    comments: Sequence[RawComment],
    records: Sequence[Mapping[str, Any]],
    policy: TimePolicy,
    secret: bytes,
    observation_ends: Mapping[tuple[int, str], datetime],
) -> dict[str, Any]:
    """Independently recompute author/course history and source horizons."""

    post_by_key = {item.raw_key: item for item in posts}
    valid_comments = [
        item
        for item in comments
        if item.post_key in post_by_key
        and item.created_at >= post_by_key[item.post_key].created_at
    ]
    comments_by_post = _comments_by_post(valid_comments)
    author_posts: dict[str, list[RawPost]] = defaultdict(list)
    course_post_times: dict[
        tuple[int, str], list[datetime]
    ] = defaultdict(list)
    course_external_times: dict[
        tuple[int, str], list[datetime]
    ] = defaultdict(list)
    first_external: dict[tuple[int, str, str], datetime | None] = {}
    for post in posts:
        author_posts[post.author_id].append(post)
        course_post_times[(post.year, post.course_id)].append(
            post.created_at
        )
        external = _external_comments(
            post, comments_by_post.get(post.raw_key, ())
        )
        first_external[post.raw_key] = (
            external[0].created_at if external else None
        )
        course_external_times[(post.year, post.course_id)].extend(
            item.created_at for item in external
        )
    for values in author_posts.values():
        values.sort(key=lambda item: (item.created_at, item.raw_key))
    for mapping in (course_post_times, course_external_times):
        for values in mapping.values():
            values.sort()

    formal_by_key = {
        str(row["snapshot_key"]): row for row in records
    }
    mismatches_by_field: Counter[str] = Counter()
    compared = 0
    for post in sorted(
        posts, key=lambda item: (item.created_at, item.raw_key)
    ):
        decision_at = post.created_at + timedelta(
            hours=policy.decision_lag_hours
        )
        outcome_end = post.created_at + timedelta(
            hours=policy.outcome_end_hours
        )
        snapshot_key = _stable_key(
            secret,
            "snapshot",
            post.year,
            post.course_id,
            post.module_id,
            _iso(decision_at),
        )
        formal = formal_by_key.get(snapshot_key)
        if formal is None:
            mismatches_by_field["snapshot_key"] += 1
            continue

        prior = [
            item
            for item in author_posts.get(post.author_id, ())
            if item.raw_key != post.raw_key
            and item.created_at < decision_at
        ]
        prior_responses = sum(
            1
            for item in prior
            if first_external[item.raw_key] is not None
            and first_external[item.raw_key] <= decision_at  # type: ignore[operator]
        )
        prior_gaps = sum(
            1
            for item in prior
            if item.created_at
            + timedelta(hours=policy.outcome_end_hours)
            <= decision_at
            and (
                first_external[item.raw_key] is None
                or first_external[item.raw_key]
                > item.created_at
                + timedelta(hours=policy.outcome_end_hours)  # type: ignore[operator]
            )
        )
        course_key = (post.year, post.course_id)
        lower = decision_at - timedelta(days=7)
        post_times = course_post_times[course_key]
        external_times = course_external_times[course_key]
        source_end = min(
            outcome_end,
            observation_ends.get(course_key, post.created_at),
        )
        expected: dict[str, int | float] = {
            "author_prior_post_count_asof": len(prior),
            "author_prior_external_response_count_asof": (
                prior_responses
            ),
            "author_prior_external_response_rate_asof": round(
                prior_responses / len(prior) if prior else 0.0,
                6,
            ),
            "author_prior_matured_gap_count_asof": prior_gaps,
            "course_post_count_prior_7d_asof": (
                bisect.bisect_right(post_times, decision_at)
                - bisect.bisect_right(post_times, lower)
            ),
            "course_external_comment_count_prior_7d_asof": (
                bisect.bisect_right(external_times, decision_at)
                - bisect.bisect_right(external_times, lower)
            ),
            "source_observation_end_after_decision_microseconds": (
                _timedelta_microseconds(source_end - decision_at)
            ),
        }
        for field, expected_value in expected.items():
            actual = formal.get(field)
            if field == "author_prior_external_response_rate_asof":
                matches = float(actual or 0) == float(expected_value)
            else:
                matches = int(actual) == int(expected_value)
            if not matches:
                mismatches_by_field[field] += 1
        compared += 1

    mismatch_count = sum(mismatches_by_field.values())
    result: dict[str, Any] = {
        "schema_version": "cmooc-history-feature-audit-v1",
        "contains_row_level_data": False,
        "status": "passed" if mismatch_count == 0 else "failed",
        "decision_count": len(records),
        "compared_row_count": compared,
        "audited_fields": list(HISTORY_FEATURE_AUDIT_FIELDS),
        "mismatch_count": mismatch_count,
        "mismatches_by_field": dict(
            sorted(mismatches_by_field.items())
        ),
        "boundary_contract": {
            "prior_post": "created_at < decision_at",
            "prior_response": "first_external_response_at <= decision_at",
            "prior_matured_gap": "outcome_end_at <= decision_at",
            "course_prior_7d": (
                "decision_at - 7d < event_at <= decision_at"
            ),
            "source_state_horizon": (
                "min(course_observation_end,outcome_end) - decision_at"
            ),
        },
    }
    result["audit_hash"] = _canonical_hash(result)
    if result["status"] != "passed":
        raise DecisionSnapshotError("HISTORY_FEATURE_AUDIT_MISMATCH")
    return result


def _structural_score(row: Mapping[str, Any]) -> float:
    response_rate = float(
        row.get("author_prior_external_response_rate_asof", 0) or 0
    )
    comment_count = int(row.get("comment_count_asof", 0) or 0)
    self_count = int(row.get("self_comment_count_asof", 0) or 0)
    depth = int(row.get("max_reply_depth_asof", 0) or 0)
    prior_gaps = int(row.get("author_prior_matured_gap_count_asof", 0) or 0)
    cognitive_need = min(24.0, depth * 1.6 + 1.0)
    urgency_need = min(
        24.0, 14.0 + (5.0 if comment_count <= 1 else 0.0) + self_count * 2.0
    )
    equity_need = min(
        24.0, (1.0 - min(max(response_rate, 0.0), 1.0)) * 18.0
        + (5.0 if comment_count <= 1 else 0.0)
    )
    history_need = min(12.0, prior_gaps * 2.0)
    return round(
        min(120.0, 50.0 + cognitive_need + urgency_need + equity_need + history_need),
        4,
    )


def evaluate_scheduler(
    records: Sequence[Mapping[str, Any]],
    *,
    budgets: Sequence[int] = DEFAULT_BUDGETS,
) -> dict[str, Any]:
    return evaluate_replay_benchmark(records, budgets=budgets)


def _score_bands(rows: Sequence[Mapping[str, Any]]) -> dict[str, str]:
    ordered = sorted(
        rows, key=lambda row: (_structural_score(row), str(row["snapshot_key"]))
    )
    result: dict[str, str] = {}
    count = len(ordered)
    for index, row in enumerate(ordered):
        fraction = (index + 0.5) / max(1, count)
        band = "low" if fraction < 1 / 3 else "mid" if fraction < 2 / 3 else "high"
        result[str(row["snapshot_key"])] = band
    return result


def _reference_strata(
    records: Sequence[Mapping[str, Any]],
) -> tuple[list[Mapping[str, Any]], dict[str, str]]:
    pool = [
        row
        for row in records
        if row.get("decision_eligibility") == "eligible_unanswered"
        and int(row.get("outcome_observable", 0) or 0) == 1
        and int(row.get("year", 0)) in (DEVELOPMENT_YEAR, HOLDOUT_YEAR)
    ]
    bands: dict[str, str] = {}
    for year in (DEVELOPMENT_YEAR, HOLDOUT_YEAR):
        cohort = [row for row in pool if int(row["year"]) == year]
        bands.update(_score_bands(cohort))
    opaque = {
        ("low", 0): "S01",
        ("low", 1): "S02",
        ("mid", 0): "S03",
        ("mid", 1): "S04",
        ("high", 0): "S05",
        ("high", 1): "S06",
    }
    stratum_by_snapshot = {
        str(row["snapshot_key"]): opaque[
            (
                bands[str(row["snapshot_key"])],
                int(row["support_gap_at_horizon"]),
            )
        ]
        for row in pool
    }
    return pool, stratum_by_snapshot


def select_reference_items(
    records: Sequence[Mapping[str, Any]], *, limit: int = DEFAULT_REFERENCE_LIMIT
) -> list[dict[str, Any]]:
    limit = max(1, min(DEFAULT_REFERENCE_LIMIT, int(limit)))
    pool, stratum_by_snapshot = _reference_strata(records)
    strata: dict[tuple[int, str], list[Mapping[str, Any]]] = defaultdict(list)
    for row in pool:
        code = stratum_by_snapshot[str(row["snapshot_key"])]
        strata[(int(row["year"]), code)].append(row)

    per_stratum = max(1, limit // 12)
    chosen: list[tuple[Mapping[str, Any], str]] = []
    chosen_keys: set[str] = set()
    chosen_authors: Counter[str] = Counter()
    for key in sorted(strata):
        candidates = sorted(
            strata[key],
            key=lambda row: (
                chosen_authors[str(row["author_key"])],
                hashlib.sha256(str(row["snapshot_key"]).encode()).hexdigest(),
            ),
        )
        for row in candidates[:per_stratum]:
            chosen.append((row, key[1]))
            chosen_keys.add(str(row["snapshot_key"]))
            chosen_authors[str(row["author_key"])] += 1

    if len(chosen) < limit:
        remaining = [row for row in pool if str(row["snapshot_key"]) not in chosen_keys]
        remaining.sort(
            key=lambda row: (
                chosen_authors[str(row["author_key"])],
                hashlib.sha256(str(row["snapshot_key"]).encode()).hexdigest(),
            )
        )
        for row in remaining[: limit - len(chosen)]:
            code = stratum_by_snapshot[str(row["snapshot_key"])]
            chosen.append((row, code))
            chosen_authors[str(row["author_key"])] += 1

    chosen = chosen[:limit]
    chosen.sort(key=lambda item: reference_blind_order_key(item[0]))
    result: list[dict[str, Any]] = []
    for index, (row, stratum) in enumerate(chosen, 1):
        result.append(
            {
                "reference_id": f"DR{index:03d}",
                "snapshot_key": row["snapshot_key"],
                "thread_key": row["thread_key"],
                "course_key": row.get("course_key", ""),
                "author_key": row["author_key"],
                "year": row["year"],
                "decision_at": row["decision_at"],
                "sampling_stratum": stratum,
                "sampling_reason": "active_structural_band_outcome_balance_author_diversity",
                "intervene": "",
                "main_function": "",
                "target_object": "",
            }
        )
    return result


def _write_csv(
    path: Path, rows: Sequence[Mapping[str, Any]], fields: Sequence[str]
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(path)


def _write_json(path: Path, value: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    temporary.replace(path)


def _read_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    integer_fields = {
        "year",
        "decision_lag_hours",
        "outcome_window_hours",
        "comment_count_asof",
        "external_comment_count_asof",
        "self_comment_count_asof",
        "unique_external_author_count_asof",
        "reply_edge_count_asof",
        "max_reply_depth_asof",
        "future_parent_reference_count_asof",
        "author_prior_post_count_asof",
        "author_prior_external_response_count_asof",
        "author_prior_matured_gap_count_asof",
        "course_post_count_prior_7d_asof",
        "course_external_comment_count_prior_7d_asof",
        "outcome_observable",
        "source_observation_end_after_decision_microseconds",
        "first_external_response_after_decision_microseconds",
        "organic_external_response_in_window",
        "external_response_count_in_window",
        "unique_external_responder_count_in_window",
        "support_gap_at_horizon",
        "asof_boundary_ok",
    }
    float_fields = {"author_prior_external_response_rate_asof"}
    for row in rows:
        for field in integer_fields:
            if field in row and row[field] != "":
                row[field] = int(row[field])
        for field in float_fields:
            if field in row and row[field] != "":
                row[field] = float(row[field])
    return rows


def _ensure_secret() -> bytes:
    path = _research_home() / "secrets" / "decision-snapshot-v1.key"
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.is_file():
        try:
            value = base64.b64decode(path.read_text(encoding="ascii").strip())
        except (OSError, ValueError) as exc:
            raise DecisionSnapshotError("INVALID_DECISION_KEY") from exc
        if len(value) != 32:
            raise DecisionSnapshotError("INVALID_DECISION_KEY")
        return value
    value = secrets.token_bytes(32)
    path.write_text(base64.b64encode(value).decode("ascii"), encoding="ascii")
    try:
        path.chmod(0o600)
    except OSError:
        pass
    return value


def _first_response_summary(
    posts: Sequence[RawPost], comments: Sequence[RawComment]
) -> dict[str, Any]:
    by_post = _comments_by_post(comments)
    result: dict[str, Any] = {}
    for year in (DEVELOPMENT_YEAR, HOLDOUT_YEAR):
        values: list[float] = []
        responded = 0
        for post in posts:
            if post.year != year:
                continue
            external = _external_comments(post, by_post.get(post.raw_key, ()))
            if external:
                responded += 1
                values.append(
                    (external[0].created_at - post.created_at).total_seconds()
                    / 3600
                )
        result[str(year)] = {
            "thread_count": sum(1 for post in posts if post.year == year),
            "externally_responded_thread_count": responded,
            "p10_hours": _quantile(values, 0.10),
            "p25_hours": _quantile(values, 0.25),
            "p50_hours": _quantile(values, 0.50),
            "p75_hours": _quantile(values, 0.75),
            "p90_hours": _quantile(values, 0.90),
            "p95_hours": _quantile(values, 0.95),
        }
    return result


def build_artifacts(
    input_root: Path, output_dir: Path
) -> tuple[list[dict[str, Any]], list[RawPost], dict[str, Any]]:
    posts, comments, source_audit = load_raw_events(input_root)
    policy = calibrate_time_policy(posts, comments)
    secret = _ensure_secret()
    observation_ends = _course_observation_ends(posts, comments)
    records = build_snapshot_records(
        posts,
        comments,
        policy,
        secret,
        observation_ends=observation_ends,
    )
    _write_csv(
        output_dir / "decision_snapshots.csv", records, FORMAL_SNAPSHOT_FIELDS
    )
    temporal_manifest = contract_manifest()
    temporal_audit = audit_temporal_rows(records)
    temporal_audit["audit_scope"] = (
        "formal stable keys and decision clocks only; dynamic builder "
        "invariants, the independent SQL oracle, and the history-feature "
        "recomputation audit check feature/outcome boundaries"
    )
    temporal_audit["event_provenance_materialized"] = False
    temporal_audit.pop("audit_hash", None)
    temporal_audit["audit_hash"] = _canonical_hash(temporal_audit)
    point_in_time_oracle = _point_in_time_oracle_audit(
        posts,
        comments,
        records,
        policy,
        secret,
    )
    history_feature_audit = _history_feature_audit(
        posts,
        comments,
        records,
        policy,
        secret,
        observation_ends,
    )
    _write_json(
        output_dir / "audit-v2" / "temporal-contract.json",
        temporal_manifest,
    )
    _write_json(
        output_dir / "audit-v2" / "decision-clock-audit.json",
        temporal_audit,
    )
    _write_json(
        output_dir / "audit-v2" / "point-in-time-oracle.json",
        point_in_time_oracle,
    )
    _write_json(
        output_dir / "audit-v2" / "history-feature-audit.json",
        history_feature_audit,
    )
    source_audit["first_external_response_distribution"] = (
        _first_response_summary(posts, comments)
    )
    source_audit["future_information_leakage_checks"] = {
        "formal_schema_records_checked": len(records),
        "asof_boundary_violation_count": sum(
            1 for row in records if int(row["asof_boundary_ok"]) != 1
        ),
        "unresolved_parent_references_excluded_from_depth": sum(
            int(row["future_parent_reference_count_asof"]) for row in records
        ),
        "final_snapshot_fields_used_as_features": [],
        "text_derived_features_used": [],
        "temporal_contract_schema_version": temporal_manifest[
            "schema_version"
        ],
        "temporal_contract_field_coverage": {
            "formal_field_count": len(FORMAL_SNAPSHOT_FIELDS),
            "contract_field_count": len(temporal_manifest["fields"]),
            "uncontracted_fields": sorted(
                set(FORMAL_SNAPSHOT_FIELDS)
                - {
                    str(item["name"])
                    for item in temporal_manifest["fields"]
                }
            ),
        },
        "point_in_time_oracle": {
            "status": point_in_time_oracle["status"],
            "backend": point_in_time_oracle["backend"],
            "decision_count": point_in_time_oracle["decision_count"],
            "event_count": point_in_time_oracle["event_count"],
            "formal_snapshot_mismatch_count": point_in_time_oracle[
                "formal_snapshot_mismatch_count"
            ],
        },
        "history_feature_audit": {
            "status": history_feature_audit["status"],
            "decision_count": history_feature_audit[
                "decision_count"
            ],
            "audited_fields": history_feature_audit[
                "audited_fields"
            ],
            "mismatch_count": history_feature_audit[
                "mismatch_count"
            ],
        },
    }
    _write_json(output_dir / "source_audit.json", source_audit)
    manifest = {
        "schema_version": SCHEMA_VERSION,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "input_root": "<local-research-data-root>",
        "output_dir": (
            "<local-research-artifact-root>/p3a/"
            "decision-snapshot-v1"
        ),
        "input_source_id": _canonical_hash(
            {"source_files": source_audit["source_files"]}
        ),
        "row_count": len(records),
        "formal_fields": list(FORMAL_SNAPSHOT_FIELDS),
        "time_policy": asdict(policy),
        "source_time_assumption": SOURCE_TIME_ASSUMPTION,
        "split_policy": {
            "development_year": DEVELOPMENT_YEAR,
            "time_holdout_year": HOLDOUT_YEAR,
        },
        "privacy": {
            "stable_keys": "HMAC-SHA256 truncated to 24 hexadecimal characters",
            "key_location": "%LOCALAPPDATA%/cmooc-research/secrets (not exported)",
            "forum_text_written": False,
            "raw_platform_ids_written_to_formal_outputs": False,
        },
        "interpretation_boundary": (
            "outcomes are observed subsequent human interaction, not effects "
            "of an intervention that did not occur"
        ),
        "temporal_contract": {
            "schema_version": temporal_manifest["schema_version"],
            "contract_hash": temporal_manifest["contract_hash"],
            "decision_clock_audit_status": temporal_audit["status"],
            "point_in_time_oracle_status": point_in_time_oracle["status"],
            "point_in_time_oracle_backend": point_in_time_oracle["backend"][
                "selected_backend"
            ],
            "history_feature_audit_status": (
                history_feature_audit["status"]
            ),
        },
    }
    _write_json(output_dir / "manifest.json", manifest)
    return records, posts, manifest


def benchmark_artifacts(
    output_dir: Path, records: Sequence[Mapping[str, Any]] | None = None
) -> dict[str, Any]:
    if records is None:
        snapshot_path = output_dir / "decision_snapshots.csv"
        if not snapshot_path.is_file():
            raise DecisionSnapshotError("DECISION_SNAPSHOTS_REQUIRED")
        records = _read_csv(snapshot_path)
    result = evaluate_scheduler(records)
    benchmark_dir = output_dir / "benchmark-v2"
    action_log = list(result.get("action_log", []))
    metrics = {
        key: value for key, value in result.items() if key != "action_log"
    }
    _write_json(benchmark_dir / "metrics.json", metrics)
    _write_csv(
        benchmark_dir / "actions.csv",
        action_log,
        BENCHMARK_ACTION_FIELDS,
    )
    snapshot_path = output_dir / "decision_snapshots.csv"
    actions_path = benchmark_dir / "actions.csv"
    _write_json(
        benchmark_dir / "replay-manifest.json",
        {
            "schema_version": BENCHMARK_SCHEMA_VERSION,
            "action_count": len(action_log),
            "action_fields": list(BENCHMARK_ACTION_FIELDS),
            "contains_forum_text": False,
            "contains_raw_platform_ids": False,
            "replay_contract": metrics["replay_contract"],
            "workload_contract": metrics["workload_contract"],
            "policy_definitions": metrics["policy_definitions"],
            "implementation_hashes": metrics[
                "implementation_hashes"
            ],
            "input_hashes_by_year": metrics["input_hashes_by_year"],
            "policy_replay_input_hash": metrics[
                "policy_replay_input_hash"
            ],
            "benchmark_input_hash": metrics["benchmark_input_hash"],
            "action_log_hash": metrics["action_log_hash"],
            "metrics_result_hash": metrics["metrics_result_hash"],
            "artifact_file_hashes": {
                "actions_csv_sha256": _sha256_file(actions_path),
                "metrics_json_sha256": _sha256_file(
                    benchmark_dir / "metrics.json"
                ),
                "decision_snapshots_csv_sha256": (
                    _sha256_file(snapshot_path)
                    if snapshot_path.is_file()
                    else None
                ),
            },
            "interpretation_boundary": metrics["interpretation_boundary"],
        },
    )
    return result


def reference_artifacts(
    output_dir: Path,
    records: Sequence[Mapping[str, Any]] | None = None,
    *,
    limit: int = DEFAULT_REFERENCE_LIMIT,
) -> list[dict[str, Any]]:
    if records is None:
        snapshot_path = output_dir / "decision_snapshots.csv"
        if not snapshot_path.is_file():
            raise DecisionSnapshotError("DECISION_SNAPSHOTS_REQUIRED")
        records = _read_csv(snapshot_path)
    reference_dir = output_dir / "reference"
    reference_path = reference_dir / "decision-reference-60.csv"
    manifest_path = reference_dir / "reference-manifest.json"
    selection_frozen = reference_path.is_file()
    existing_manifest: dict[str, Any] = {}
    if selection_frozen:
        with reference_path.open(
            "r", encoding="utf-8-sig", newline=""
        ) as handle:
            reader = csv.DictReader(handle)
            if tuple(reader.fieldnames or ()) != REFERENCE_FIELDS:
                raise DecisionSnapshotError("REFERENCE_SCHEMA_MISMATCH")
            selected = list(reader)
        for item in selected:
            if item.get("year") not in ("", None):
                item["year"] = int(item["year"])
        if not selected or len(selected) > DEFAULT_REFERENCE_LIMIT:
            raise DecisionSnapshotError("INVALID_FROZEN_REFERENCE_SIZE")
        if manifest_path.is_file():
            try:
                loaded_manifest = json.loads(
                    manifest_path.read_text(encoding="utf-8")
                )
            except (
                OSError,
                UnicodeError,
                json.JSONDecodeError,
            ) as exc:
                raise DecisionSnapshotError(
                    "REFERENCE_MANIFEST_INVALID"
                ) from exc
            if not isinstance(loaded_manifest, dict):
                raise DecisionSnapshotError(
                    "REFERENCE_MANIFEST_INVALID"
                )
            existing_manifest = loaded_manifest
    else:
        selected = select_reference_items(records, limit=limit)
        _write_csv(reference_path, selected, REFERENCE_FIELDS)

    pool, stratum_by_snapshot = _reference_strata(records)
    pool_by_key = {str(row["snapshot_key"]): row for row in pool}
    allowed_labels = {
        "intervene": {"", "yes", "no"},
        "main_function": {
            "",
            "resource_connection",
            "cognitive_conflict",
            "emotional_support",
            "none",
        },
        "target_object": {
            "",
            "thread_author",
            "specific_participant",
            "thread_participants",
        },
    }
    seen_snapshots: set[str] = set()
    for item in selected:
        snapshot_key = str(item["snapshot_key"])
        if snapshot_key in seen_snapshots:
            raise DecisionSnapshotError(
                f"DUPLICATE_FROZEN_REFERENCE:{snapshot_key}"
            )
        seen_snapshots.add(snapshot_key)
        if snapshot_key not in pool_by_key:
            raise DecisionSnapshotError(
                f"FROZEN_REFERENCE_NOT_IN_SNAPSHOT:{snapshot_key}"
            )
        expected_stratum = stratum_by_snapshot[snapshot_key]
        if str(item["sampling_stratum"]) != expected_stratum:
            raise DecisionSnapshotError(
                f"FROZEN_REFERENCE_STRATUM_MISMATCH:{snapshot_key}"
            )
        for field, allowed in allowed_labels.items():
            if str(item.get(field, "")) not in allowed:
                raise DecisionSnapshotError(
                    f"INVALID_REFERENCE_LABEL:{field}:{snapshot_key}"
                )

    legacy_blind_order_migrated = bool(
        existing_manifest.get("legacy_blind_order_migrated", False)
    )
    if selection_frozen:
        if existing_manifest and (
            existing_manifest.get("schema_version")
            != REFERENCE_SCHEMA_VERSION
            or existing_manifest.get("item_count") != len(selected)
            or existing_manifest.get("selection_frozen") is not True
        ):
            raise DecisionSnapshotError(
                "REFERENCE_MANIFEST_INVALID"
            )
        try:
            has_integrity = has_complete_reference_integrity(
                existing_manifest
            )
        except ReferenceIntegrityError as exc:
            raise DecisionSnapshotError(exc.code) from exc
        if has_integrity:
            try:
                verify_reference_integrity(
                    existing_manifest, selected
                )
            except ReferenceIntegrityError as exc:
                raise DecisionSnapshotError(exc.code) from exc
        else:
            canonical = sorted(
                selected, key=reference_blind_order_key
            )
            canonical = [
                {
                    **item,
                    "reference_id": f"DR{index:03d}",
                }
                for index, item in enumerate(canonical, start=1)
            ]
            order_changed = any(
                left.get("reference_id") != right.get("reference_id")
                or left.get("snapshot_key") != right.get("snapshot_key")
                for left, right in zip(selected, canonical)
            )
            if order_changed:
                if any(
                    str(item.get(field, "")).strip()
                    for item in selected
                    for field in REFERENCE_LABEL_FIELDS
                ):
                    raise DecisionSnapshotError(
                        "LABELLED_REFERENCE_ORDER_MIGRATION_FORBIDDEN"
                    )
                selected = canonical
                _write_csv(
                    reference_path, selected, REFERENCE_FIELDS
                )
                legacy_blind_order_migrated = True

    integrity = reference_integrity(selected)
    population_counts = Counter(
        (str(row["year"]), stratum_by_snapshot[str(row["snapshot_key"])])
        for row in pool
    )
    counts = Counter(
        (str(item["year"]), str(item["sampling_stratum"])) for item in selected
    )
    population_named = {
        f"{year}:{stratum}": count
        for (year, stratum), count in sorted(population_counts.items())
    }
    selected_named = {
        f"{year}:{stratum}": counts.get((year, stratum), 0)
        for (year, stratum) in sorted(population_counts)
    }
    manifest = {
        "schema_version": REFERENCE_SCHEMA_VERSION,
        "item_count": len(selected),
        "maximum_item_count": DEFAULT_REFERENCE_LIMIT,
        "selection_frozen": True,
        **integrity,
        "legacy_blind_order_migrated": (
            legacy_blind_order_migrated
        ),
        "label_fields": list(REFERENCE_LABEL_FIELDS),
        "allowed_labels": {
            "intervene": ["yes", "no"],
            "main_function": [
                "resource_connection",
                "cognitive_conflict",
                "emotional_support",
                "none",
            ],
            "target_object": [
                "thread_author",
                "specific_participant",
                "thread_participants",
            ],
        },
        "sampling": (
            "30 per year when available; balanced over three structural-score "
            "bands and two future observational outcomes; author-diverse "
            "deterministic active selection. The coder sheet hides outcomes."
        ),
        "stratum_counts": selected_named,
        "stratum_population_counts": population_named,
        "stratum_inclusion_probabilities": {
            key: selected_named[key] / population
            for key, population in population_named.items()
        },
        "instruction": (
            "Complete only intervene, main_function, and target_object. Stop "
            "after this first batch; do not expand coding automatically."
        ),
    }
    _write_json(manifest_path, manifest)
    return selected


def status(output_dir: Path) -> dict[str, Any]:
    paths = {
        "manifest": output_dir / "manifest.json",
        "source_audit": output_dir / "source_audit.json",
        "snapshots": output_dir / "decision_snapshots.csv",
        "benchmark_v2": output_dir / "benchmark-v2" / "metrics.json",
        "benchmark_v2_actions": output_dir / "benchmark-v2" / "actions.csv",
        "temporal_contract_v2": (
            output_dir / "audit-v2" / "temporal-contract.json"
        ),
        "decision_clock_audit_v2": (
            output_dir / "audit-v2" / "decision-clock-audit.json"
        ),
        "point_in_time_oracle_v2": (
            output_dir / "audit-v2" / "point-in-time-oracle.json"
        ),
        "history_feature_audit_v1": (
            output_dir / "audit-v2" / "history-feature-audit.json"
        ),
        "reference": output_dir / "reference" / "decision-reference-60.csv",
        "reference_manifest": output_dir / "reference" / "reference-manifest.json",
    }
    return {
        "ok": True,
        "command": "status",
        "network_used": False,
        "output_dir": str(output_dir),
        "artifacts": {name: path.is_file() for name, path in paths.items()},
    }


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Local-only decision snapshot and scheduler benchmark"
    )
    subparsers = parser.add_subparsers(dest="command", required=True)
    for command in ("status", "audit", "build", "benchmark", "reference", "all"):
        child = subparsers.add_parser(command)
        child.add_argument("--input-root")
        child.add_argument("--output-dir")
        child.add_argument("--json", action="store_true")
        if command in ("reference", "all"):
            child.add_argument(
                "--limit", type=int, default=DEFAULT_REFERENCE_LIMIT
            )
    return parser


def _emit(value: Mapping[str, Any], json_mode: bool) -> None:
    if json_mode:
        print(json.dumps(value, ensure_ascii=False, sort_keys=True))
        return
    if value.get("ok") is False:
        print(f"决策快照流程未继续：{value.get('error')}")
        return
    print(f"决策快照流程：{value.get('command')} 已完成")
    if value.get("output_dir"):
        print(f"本地产物：{value['output_dir']}")


def main(argv: list[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    json_mode = bool(args.json)
    try:
        input_root = (
            Path(args.input_root).expanduser().resolve()
            if args.input_root
            else _default_input_root().resolve()
        )
        output_dir = _safe_output_dir(args.output_dir)
        if args.command == "status":
            result = status(output_dir)
        elif args.command == "audit":
            posts, comments, source_audit = load_raw_events(input_root)
            policy = calibrate_time_policy(posts, comments)
            result = {
                "ok": True,
                "command": "audit",
                "network_used": False,
                "input_root": str(input_root),
                "source_audit": source_audit,
                "first_external_response_distribution": _first_response_summary(
                    posts, comments
                ),
                "time_policy": asdict(policy),
                "writes_performed": False,
            }
        elif args.command == "build":
            records, _, manifest = build_artifacts(input_root, output_dir)
            result = {
                "ok": True,
                "command": "build",
                "network_used": False,
                "output_dir": str(output_dir),
                "row_count": len(records),
                "time_policy": manifest["time_policy"],
            }
        elif args.command == "benchmark":
            benchmark = benchmark_artifacts(output_dir)
            result = {
                "ok": True,
                "command": "benchmark",
                "network_used": False,
                "output_dir": str(output_dir),
                "metric_rows": len(benchmark["by_year"]),
            }
        elif args.command == "reference":
            selected = reference_artifacts(
                output_dir, limit=int(args.limit)
            )
            result = {
                "ok": True,
                "command": "reference",
                "network_used": False,
                "output_dir": str(output_dir),
                "item_count": len(selected),
            }
        else:
            records, _, manifest = build_artifacts(input_root, output_dir)
            benchmark = benchmark_artifacts(output_dir, records)
            selected = reference_artifacts(
                output_dir, records, limit=int(args.limit)
            )
            result = {
                "ok": True,
                "command": "all",
                "network_used": False,
                "output_dir": str(output_dir),
                "row_count": len(records),
                "metric_rows": len(benchmark["by_year"]),
                "reference_count": len(selected),
                "time_policy": manifest["time_policy"],
            }
        _emit(result, json_mode)
        return 0
    except (DecisionSnapshotError, OSError, ValueError) as exc:
        result = {
            "ok": False,
            "command": args.command,
            "network_used": False,
            "error": str(exc),
        }
        _emit(result, json_mode)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
