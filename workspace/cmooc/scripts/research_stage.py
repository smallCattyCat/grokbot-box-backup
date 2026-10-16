#!/usr/bin/env python3
"""Local-only audit and staging for the historical cMOOC forum export.

This module deliberately does not import the application, Prisma, or any LLM
client. Audit is read-only. Stage writes a new SQLite database outside the
synced project tree only after all explicit safety gates pass.
"""

from __future__ import annotations

import argparse
import base64
import binascii
import hashlib
import hmac
import importlib.metadata
import json
import os
import platform
import re
import socket
import sqlite3
import sys
import tempfile
import uuid
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable, Iterator, Mapping, Sequence
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError


SCHEMA_VERSION = "p1c-role-function-staging-v1"
MANIFEST_VERSION = "p1c-source-audit-v1"
HMAC_ENV = "CMOOC_RESEARCH_HMAC_KEY_B64"
HMAC_KEY_ID_ENV = "CMOOC_RESEARCH_HMAC_KEY_ID"
LOCAL_ONLY_ENV = "CMOOC_RESEARCH_LOCAL_ONLY"
OUTPUT_ROOT_ENV = "CMOOC_RESEARCH_OUTPUT_ROOT"

MAX_COMPRESSED_BYTES = 32 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
MAX_TOTAL_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
MAX_ZIP_ENTRIES = 64
MAX_COMPRESSION_RATIO = 20.0
MAX_ROWS = 200_000
MAX_COLUMNS = 128
MAX_CELL_LENGTH = 32_767

VALID_COURSE_ROLES = {"learner", "helper", "creator", "unknown"}
VALID_SOURCE_KINDS = {"human", "companion_generated", "system", "test", "unknown"}
ROLE_MAP = {"LEARNER": "learner", "HELPER": "helper", "CREATOR": "creator"}
STRUCTURE_SIGNAL_VERSION = "structure-v1"
URL_PATTERN = re.compile(r"https?://[^\s<>\"']+", re.I)
BLOCK_TAGS = {
    "address",
    "article",
    "aside",
    "blockquote",
    "br",
    "div",
    "h1",
    "h2",
    "h3",
    "h4",
    "h5",
    "h6",
    "li",
    "p",
    "pre",
    "section",
    "table",
    "tr",
}

PII_PATTERNS: Mapping[str, re.Pattern[str]] = {
    "email": re.compile(r"(?<![\w.+-])[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}(?![\w.-])"),
    "phone": re.compile(r"(?<!\d)(?:\+?86[- ]?)?1[3-9]\d{9}(?!\d)"),
    "identity_number": re.compile(r"(?<!\d)\d{17}[\dXx](?!\d)"),
    "contact_handle": re.compile(r"(?:微信|vx|wechat|qq)\s*[:：]?\s*[A-Za-z0-9_-]{5,}", re.I),
    "credential_or_token": re.compile(r"(?:api[_-]?key|access[_-]?token|password|密码)\s*[:=：]\s*\S+", re.I),
    "url_with_query": re.compile(r"https?://[^\s?]+\?[^\s]+", re.I),
}


class ResearchStageError(RuntimeError):
    """Expected, sanitized command failure."""

    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


class _StructureTextParser(HTMLParser):
    """Extract privacy-safe structure signals without retaining source text."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        normalized = tag.lower()
        if normalized in {"script", "style"}:
            self._ignored_depth += 1
            return
        if not self._ignored_depth and normalized in BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        normalized = tag.lower()
        if normalized in {"script", "style"} and self._ignored_depth:
            self._ignored_depth -= 1
            return
        if not self._ignored_depth and normalized in BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._ignored_depth:
            self.parts.append(data)


def _content_structure(content: str) -> tuple[int, int, int]:
    parser = _StructureTextParser()
    try:
        parser.feed(content)
        parser.close()
        plain_text = "".join(parser.parts)
    except Exception:
        plain_text = re.sub(r"<[^>]*>", " ", content)
    normalized_char_count = len(re.sub(r"\s+", "", plain_text, flags=re.UNICODE))
    paragraph_count = sum(
        1 for part in re.split(r"(?:\r?\n)+", plain_text) if part.strip()
    )
    url_count = len(URL_PATTERN.findall(content))
    return normalized_char_count, paragraph_count, url_count


@dataclass(frozen=True)
class FileSpec:
    logical_name: str
    relative_path: str
    entity_type: str
    year: int
    expected_headers: tuple[str, ...]
    read_columns: tuple[str, ...]


@dataclass(frozen=True)
class FileAudit:
    logical_name: str
    relative_path: str
    sha256: str
    byte_size: int
    worksheet: str
    business_rows: int
    accepted_rows: int = 0
    quarantined_rows: int = 0

    def manifest_dict(self) -> dict[str, Any]:
        return {
            "logical_name": self.logical_name,
            "relative_path": self.relative_path.replace("\\", "/"),
            "sha256": self.sha256,
            "byte_size": self.byte_size,
            "worksheet": self.worksheet,
            "business_rows": self.business_rows,
            "accepted_rows": self.accepted_rows,
            "quarantined_rows": self.quarantined_rows,
        }


@dataclass(frozen=True)
class CommentLink:
    logical_name: str
    course_id: str | None
    comment_id: str | None
    module_id: str | None
    actor_id: str | None
    is_reply: bool
    parent_comment_id: str | None


@dataclass
class AuditState:
    course_ids: set[str] = field(default_factory=set)
    memberships: dict[tuple[str, str], str] = field(default_factory=dict)
    post_keys: set[tuple[str, str]] = field(default_factory=set)
    post_ids: set[tuple[str, str]] = field(default_factory=set)
    comment_ids: set[tuple[str, str]] = field(default_factory=set)
    canonical_comment_ids: set[tuple[str, str]] = field(default_factory=set)
    comments: list[CommentLink] = field(default_factory=list)
    quarantine_counts: Counter[str] = field(default_factory=Counter)
    pii_candidate_counts: Counter[str] = field(default_factory=Counter)
    role_counts: Counter[str] = field(default_factory=Counter)
    timestamp_parseable: int = 0
    timestamp_unparseable: int = 0
    timestamp_min: datetime | None = None
    timestamp_max: datetime | None = None
    empty_post_content: int = 0
    possible_truncation: int = 0
    duplicate_event_ids: int = 0
    business_rows_by_file: Counter[str] = field(default_factory=Counter)
    accepted_rows_by_file: Counter[str] = field(default_factory=Counter)


def _deny_network() -> None:
    """Make accidental external calls fail even when application env says real."""

    if os.environ.get(LOCAL_ONLY_ENV, "1") != "1":
        raise ResearchStageError("LOCAL_ONLY_REQUIRED")

    def denied(*_: Any, **__: Any) -> Any:
        raise ResearchStageError("NETWORK_ACCESS_DENIED")

    base_socket = socket.socket

    class LocalOnlySocket(base_socket):  # type: ignore[misc,valid-type]
        def connect(self, *_: Any, **__: Any) -> Any:
            return denied()

        def connect_ex(self, *_: Any, **__: Any) -> Any:
            return denied()

    socket.socket = LocalOnlySocket  # type: ignore[assignment]
    socket.create_connection = denied  # type: ignore[assignment]


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[1]


def _source_schema_path() -> Path:
    return _repo_root() / "research" / "source-schema-v1.json"


def _migration_paths() -> tuple[Path, ...]:
    paths = tuple(sorted((_repo_root() / "research" / "migrations").glob("*.sql")))
    if not paths:
        raise ResearchStageError("MIGRATION_UNAVAILABLE")
    return paths


def _migration_bundle_sha256(paths: Sequence[Path]) -> str:
    digest = hashlib.sha256()
    for path in paths:
        digest.update(path.name.encode("utf-8"))
        digest.update(b"\0")
        digest.update(path.read_bytes())
        digest.update(b"\0")
    return digest.hexdigest()


def _load_specs() -> tuple[str, list[FileSpec]]:
    try:
        raw = json.loads(_source_schema_path().read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ResearchStageError("SOURCE_SCHEMA_UNAVAILABLE") from exc
    try:
        specs = [
            FileSpec(
                logical_name=item["logical_name"],
                relative_path=item["relative_path"],
                entity_type=item["entity_type"],
                year=int(item["year"]),
                expected_headers=tuple(item["expected_headers"]),
                read_columns=tuple(item["read_columns"]),
            )
            for item in raw["files"]
        ]
    except (KeyError, TypeError, ValueError) as exc:
        raise ResearchStageError("SOURCE_SCHEMA_INVALID") from exc
    if raw.get("schema_version") != SCHEMA_VERSION or len(specs) != 8:
        raise ResearchStageError("SOURCE_SCHEMA_INVALID")
    return str(raw.get("source_subdir", "")), specs


def _is_unc(path: Path) -> bool:
    return str(path).startswith("\\\\")


def _contains(parent: Path, child: Path) -> bool:
    try:
        child.relative_to(parent)
        return True
    except ValueError:
        return False


def _resolve_input_root(value: str) -> Path:
    raw = Path(value).expanduser()
    if _is_unc(raw):
        raise ResearchStageError("UNC_INPUT_FORBIDDEN")
    try:
        root = raw.resolve(strict=True)
    except OSError as exc:
        raise ResearchStageError("INPUT_ROOT_UNAVAILABLE") from exc
    if not root.is_dir():
        raise ResearchStageError("INPUT_ROOT_UNAVAILABLE")
    return root


def _resolve_source_files(input_root: Path, specs: Sequence[FileSpec]) -> dict[str, Path]:
    result: dict[str, Path] = {}
    for spec in specs:
        candidate = input_root.joinpath(*Path(spec.relative_path).parts)
        try:
            resolved = candidate.resolve(strict=True)
        except OSError as exc:
            raise ResearchStageError("SOURCE_FILE_MISSING") from exc
        if not _contains(input_root, resolved) or not resolved.is_file():
            raise ResearchStageError("SOURCE_PATH_ESCAPE")
        result[spec.logical_name] = resolved
    return result


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _preflight_xlsx(path: Path) -> None:
    if path.stat().st_size > MAX_COMPRESSED_BYTES:
        raise ResearchStageError("XLSX_COMPRESSED_LIMIT")
    try:
        with zipfile.ZipFile(path, "r") as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ZIP_ENTRIES:
                raise ResearchStageError("XLSX_ENTRY_LIMIT")
            total_uncompressed = 0
            for info in infos:
                name = info.filename.replace("\\", "/")
                parts = Path(name).parts
                if name.startswith("/") or ".." in parts:
                    raise ResearchStageError("XLSX_PATH_ESCAPE")
                total_uncompressed += info.file_size
                ratio = info.file_size / max(info.compress_size, 1)
                if ratio > MAX_COMPRESSION_RATIO:
                    raise ResearchStageError("XLSX_COMPRESSION_RATIO")
                lower = name.lower()
                if (
                    "vbaproject.bin" in lower
                    or lower.startswith("xl/externallinks/")
                    or lower.startswith("xl/embeddings/")
                ):
                    raise ResearchStageError("XLSX_ACTIVE_CONTENT")
            if total_uncompressed > MAX_UNCOMPRESSED_BYTES:
                raise ResearchStageError("XLSX_UNCOMPRESSED_LIMIT")
    except zipfile.BadZipFile as exc:
        raise ResearchStageError("XLSX_INVALID_CONTAINER") from exc


def _import_openpyxl() -> Any:
    try:
        import openpyxl  # type: ignore
    except ImportError as exc:
        raise ResearchStageError("OPENPYXL_REQUIRED") from exc
    return openpyxl


def _header_indices(headers: Sequence[Any], expected: Sequence[str]) -> dict[str, int]:
    normalized = tuple("" if value is None else str(value).strip() for value in headers)
    if normalized != tuple(expected):
        raise ResearchStageError("SOURCE_HEADER_MISMATCH")
    indices: dict[str, int] = {}
    for index, name in enumerate(normalized):
        if name not in indices:
            indices[name] = index
    return indices


def _inspect_workbook(path: Path, spec: FileSpec) -> tuple[str, int]:
    openpyxl = _import_openpyxl()
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        if len(workbook.sheetnames) != 1:
            raise ResearchStageError("XLSX_SHEET_COUNT")
        worksheet = workbook[workbook.sheetnames[0]]
        if worksheet.max_row > MAX_ROWS or worksheet.max_column > MAX_COLUMNS:
            raise ResearchStageError("XLSX_DIMENSION_LIMIT")
        header_cells = next(worksheet.iter_rows(min_row=1, max_row=1))
        _header_indices([cell.value for cell in header_cells], spec.expected_headers)
        for cells in worksheet.iter_rows(min_row=1, max_row=min(2, worksheet.max_row)):
            if any(cell.data_type == "f" for cell in cells):
                raise ResearchStageError("XLSX_FORMULA_FORBIDDEN")
            if any(
                isinstance(cell.value, str) and len(cell.value) > MAX_CELL_LENGTH
                for cell in cells
            ):
                raise ResearchStageError("XLSX_CELL_LIMIT")
        return "first_worksheet", max(worksheet.max_row - 2, 0)
    except ResearchStageError:
        raise
    except Exception as exc:
        raise ResearchStageError("XLSX_READ_FAILED") from exc
    finally:
        if "workbook" in locals():
            workbook.close()


def _iter_rows(path: Path, spec: FileSpec) -> Iterator[dict[str, Any]]:
    openpyxl = _import_openpyxl()
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        worksheet = workbook[workbook.sheetnames[0]]
        header = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        indices = _header_indices(header, spec.expected_headers)
        missing = [name for name in spec.read_columns if name not in indices]
        if missing:
            raise ResearchStageError("SOURCE_REQUIRED_COLUMN_MISSING")
        for cells in worksheet.iter_rows(min_row=3):
            if any(cell.data_type == "f" for cell in cells):
                raise ResearchStageError("XLSX_FORMULA_FORBIDDEN")
            values = tuple(cell.value for cell in cells)
            if any(isinstance(value, str) and len(value) > MAX_CELL_LENGTH for value in values):
                raise ResearchStageError("XLSX_CELL_LIMIT")
            if not any(value not in (None, "") for value in values):
                continue
            yield {name: values[indices[name]] for name in spec.read_columns}
    except ResearchStageError:
        raise
    except Exception as exc:
        raise ResearchStageError("XLSX_READ_FAILED") from exc
    finally:
        if "workbook" in locals():
            workbook.close()


def _source_id(value: Any) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        text = str(value)
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value).strip()
    return None if text in {"", "0"} else text


def _foreign_id(value: Any) -> str | None:
    """Normalize platform foreign-key sentinels without changing normal text/metrics."""

    normalized = _source_id(value)
    return None if normalized == "-1" else normalized


def _text(value: Any) -> str:
    return "" if value is None else str(value)


def _is_reply(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return int(value) != 0
    return str(value).strip().lower() not in {"", "0", "false", "no"}


def _parse_naive_timestamp(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = _text(value).strip()
    if not text:
        return None
    candidates = (text, text.replace("/", "-"))
    for candidate in candidates:
        try:
            parsed = datetime.fromisoformat(candidate)
            return parsed.replace(tzinfo=None)
        except ValueError:
            pass
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _timestamp_raw(value: Any) -> str:
    parsed = _parse_naive_timestamp(value)
    if parsed is None:
        return ""
    return parsed.isoformat(sep=" ", timespec="seconds")


def _to_utc(value: Any, source_timezone: ZoneInfo) -> str | None:
    parsed = _parse_naive_timestamp(value)
    if parsed is None:
        return None
    return parsed.replace(tzinfo=source_timezone).astimezone(timezone.utc).isoformat()


def _scan_pii(text: str, target: Counter[str]) -> None:
    for name, pattern in PII_PATTERNS.items():
        if pattern.search(text):
            target[name] += 1


def _update_time_range(state: AuditState, value: Any) -> None:
    parsed = _parse_naive_timestamp(value)
    if parsed is None:
        state.timestamp_unparseable += 1
        return
    state.timestamp_parseable += 1
    if state.timestamp_min is None or parsed < state.timestamp_min:
        state.timestamp_min = parsed
    if state.timestamp_max is None or parsed > state.timestamp_max:
        state.timestamp_max = parsed


def _suppress_small_counts(counts: Mapping[str, int]) -> dict[str, int | str]:
    return {
        key: (0 if value == 0 else "<5" if value < 5 else value)
        for key, value in sorted(counts.items())
    }


def _compute_canonical_comments(state: AuditState) -> None:
    candidates: dict[tuple[str, str], CommentLink] = {}
    for link in state.comments:
        if not link.course_id or not link.comment_id or not link.module_id or not link.actor_id:
            state.quarantine_counts["MISSING_REQUIRED_ID"] += 1
            continue
        key = (link.course_id, link.comment_id)
        if (link.course_id, link.module_id) not in state.post_keys:
            state.quarantine_counts["ORPHAN_POST"] += 1
            continue
        candidates[key] = link

    canonical: set[tuple[str, str]] = {
        key for key, link in candidates.items() if not link.is_reply
    }
    pending = {key for key, link in candidates.items() if link.is_reply}
    changed = True
    while changed:
        changed = False
        for key in list(pending):
            link = candidates[key]
            parent_key = (
                (link.course_id, link.parent_comment_id)
                if link.course_id and link.parent_comment_id
                else None
            )
            if parent_key and parent_key in canonical:
                canonical.add(key)
                pending.remove(key)
                changed = True
    state.canonical_comment_ids = canonical
    for key in pending:
        link = candidates[key]
        parent_key = (
            (link.course_id, link.parent_comment_id)
            if link.course_id and link.parent_comment_id
            else None
        )
        reason = "BROKEN_PARENT" if not parent_key or parent_key not in state.comment_ids else "PARENT_NOT_CANONICAL"
        state.quarantine_counts[reason] += 1
    for link in state.comments:
        if link.course_id and link.comment_id and (link.course_id, link.comment_id) in canonical:
            state.accepted_rows_by_file[link.logical_name] += 1


def _dataset_fingerprint(
    dataset_version: str, files: Sequence[FileAudit], aggregate: Mapping[str, Any]
) -> str:
    canonical = {
        "schema_version": SCHEMA_VERSION,
        "dataset_version": dataset_version,
        "files": [item.manifest_dict() for item in sorted(files, key=lambda x: x.logical_name)],
        "source_counts": aggregate.get("source_counts", {}),
    }
    encoded = json.dumps(canonical, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def _tooling_hashes() -> dict[str, str]:
    migration_paths = _migration_paths()
    tooling = {
        "importer_sha256": _sha256_file(Path(__file__).resolve()),
        "source_schema_sha256": _sha256_file(_source_schema_path()),
        "migration_sha256": _migration_bundle_sha256(migration_paths),
        "python_version": platform.python_version(),
        "openpyxl_version": str(_import_openpyxl().__version__),
        "et_xmlfile_version": importlib.metadata.version("et_xmlfile"),
        "tzdata_version": importlib.metadata.version("tzdata"),
        "sqlite_version": sqlite3.sqlite_version,
    }
    encoded = json.dumps(tooling, sort_keys=True, separators=(",", ":")).encode("utf-8")
    tooling["tooling_fingerprint"] = hashlib.sha256(encoded).hexdigest()
    return tooling


def audit_dataset(input_root_value: str, dataset_version: str) -> tuple[dict[str, Any], AuditState]:
    if not re.fullmatch(r"[A-Za-z0-9._-]{1,128}", dataset_version.strip()):
        raise ResearchStageError("DATASET_VERSION_REQUIRED")
    tooling_before = _tooling_hashes()
    _, specs = _load_specs()
    input_root = _resolve_input_root(input_root_value)
    source_paths = _resolve_source_files(input_root, specs)

    file_audits: list[FileAudit] = []
    total_uncompressed = 0
    before_hashes: dict[str, str] = {}
    for spec in specs:
        path = source_paths[spec.logical_name]
        _preflight_xlsx(path)
        with zipfile.ZipFile(path, "r") as archive:
            total_uncompressed += sum(item.file_size for item in archive.infolist())
        if total_uncompressed > MAX_TOTAL_UNCOMPRESSED_BYTES:
            raise ResearchStageError("XLSX_TOTAL_UNCOMPRESSED_LIMIT")
        sha256 = _sha256_file(path)
        before_hashes[spec.logical_name] = sha256
        worksheet, business_rows = _inspect_workbook(path, spec)
        file_audits.append(
            FileAudit(
                logical_name=spec.logical_name,
                relative_path=spec.relative_path,
                sha256=sha256,
                byte_size=path.stat().st_size,
                worksheet=worksheet,
                business_rows=business_rows,
            )
        )

    state = AuditState()
    seen_course_rows: set[str] = set()
    seen_membership_rows: set[tuple[str, str]] = set()
    seen_event_ids: set[tuple[str, str]] = set()

    for spec in specs:
        path = source_paths[spec.logical_name]
        for row in _iter_rows(path, spec):
            state.business_rows_by_file[spec.logical_name] += 1
            if spec.entity_type == "course":
                course_id = _source_id(row.get("course_id"))
                if not course_id:
                    state.quarantine_counts["MISSING_COURSE_ID"] += 1
                    continue
                if course_id in seen_course_rows:
                    state.quarantine_counts["DUPLICATE_COURSE_ID"] += 1
                    continue
                seen_course_rows.add(course_id)
                state.course_ids.add(course_id)
                state.accepted_rows_by_file[spec.logical_name] += 1

            elif spec.entity_type == "course_membership":
                course_id = _source_id(row.get("course_id"))
                actor_id = _source_id(row.get("user_id"))
                role_raw = _text(row.get("member_type")).strip().upper()
                if not course_id or not actor_id:
                    state.quarantine_counts["MISSING_ACTOR_ID"] += 1
                    continue
                role = ROLE_MAP.get(role_raw, "unknown")
                if role == "unknown":
                    state.quarantine_counts["UNKNOWN_COURSE_ROLE"] += 1
                key = (course_id, actor_id)
                if key in seen_membership_rows:
                    state.quarantine_counts["DUPLICATE_MEMBERSHIP"] += 1
                    continue
                seen_membership_rows.add(key)
                state.memberships[key] = role
                state.role_counts[role] += 1
                state.accepted_rows_by_file[spec.logical_name] += 1

            elif spec.entity_type == "post":
                course_id = _source_id(row.get("course_id"))
                module_id = _source_id(row.get("module_id"))
                actor_id = _source_id(row.get("created_by"))
                if not course_id or not module_id or not actor_id:
                    state.quarantine_counts["MISSING_REQUIRED_ID"] += 1
                    continue
                event_key = ("post", module_id)
                if event_key in seen_event_ids:
                    state.duplicate_event_ids += 1
                    state.quarantine_counts["DUPLICATE_EVENT_ID"] += 1
                    continue
                seen_event_ids.add(event_key)
                state.post_keys.add((course_id, module_id))
                state.post_ids.add((course_id, module_id))
                content = _text(row.get("module_content"))
                title = _text(row.get("module_title"))
                if not content.strip():
                    state.empty_post_content += 1
                if max(len(content), len(title)) >= MAX_CELL_LENGTH:
                    state.possible_truncation += 1
                _scan_pii(title, state.pii_candidate_counts)
                _scan_pii(content, state.pii_candidate_counts)
                _update_time_range(state, row.get("created_time"))
                state.accepted_rows_by_file[spec.logical_name] += 1

            elif spec.entity_type == "reply":
                course_id = _source_id(row.get("course_id"))
                comment_id = _source_id(row.get("comment_id"))
                module_id = _source_id(row.get("module_id"))
                actor_id = _source_id(row.get("created_by"))
                is_reply = _is_reply(row.get("comment_reply"))
                child_id = _foreign_id(row.get("comment_reply_child_id"))
                root_id = _foreign_id(row.get("comment_reply_id"))
                parent_id = child_id or root_id if is_reply else None
                if course_id and comment_id:
                    key = (course_id, comment_id)
                    if key in state.comment_ids:
                        state.duplicate_event_ids += 1
                        state.quarantine_counts["DUPLICATE_EVENT_ID"] += 1
                    state.comment_ids.add(key)
                state.comments.append(
                    CommentLink(
                        logical_name=spec.logical_name,
                        course_id=course_id,
                        comment_id=comment_id,
                        module_id=module_id,
                        actor_id=actor_id,
                        is_reply=is_reply,
                        parent_comment_id=parent_id,
                    )
                )
                content = _text(row.get("comment_content"))
                if len(content) >= MAX_CELL_LENGTH:
                    state.possible_truncation += 1
                _scan_pii(content, state.pii_candidate_counts)
                _update_time_range(state, row.get("created_time"))

    _compute_canonical_comments(state)

    source_counts = {
        "courses": sum(
            state.business_rows_by_file[spec.logical_name]
            for spec in specs
            if spec.entity_type == "course"
        ),
        "memberships": sum(
            state.business_rows_by_file[spec.logical_name]
            for spec in specs
            if spec.entity_type == "course_membership"
        ),
        "posts": sum(
            state.business_rows_by_file[spec.logical_name]
            for spec in specs
            if spec.entity_type == "post"
        ),
        "comments": sum(
            state.business_rows_by_file[spec.logical_name]
            for spec in specs
            if spec.entity_type == "reply"
        ),
    }
    accepted_counts = {
        "courses": len(state.course_ids),
        "memberships": len(state.memberships),
        "posts": len(state.post_ids),
        "comments": len(state.canonical_comment_ids),
    }

    updated_files: list[FileAudit] = []
    for item in file_audits:
        business_rows = state.business_rows_by_file[item.logical_name]
        accepted = state.accepted_rows_by_file[item.logical_name]
        updated_files.append(
            FileAudit(
                logical_name=item.logical_name,
                relative_path=item.relative_path,
                sha256=item.sha256,
                byte_size=item.byte_size,
                worksheet=item.worksheet,
                business_rows=business_rows,
                accepted_rows=accepted,
                quarantined_rows=max(business_rows - accepted, 0),
            )
        )

    for spec in specs:
        if _sha256_file(source_paths[spec.logical_name]) != before_hashes[spec.logical_name]:
            raise ResearchStageError("SOURCE_CHANGED_DURING_AUDIT")

    aggregate = {
        "source_counts": source_counts,
        "accepted_counts": accepted_counts,
        "forum_events": source_counts["posts"] + source_counts["comments"],
        "canonical_forum_events": accepted_counts["posts"] + accepted_counts["comments"],
        "quarantine_counts": dict(sorted(state.quarantine_counts.items())),
        "role_counts": dict(sorted(state.role_counts.items())),
        "timestamp": {
            "parseable": state.timestamp_parseable,
            "unparseable": state.timestamp_unparseable,
            "precision": "second",
            "local_min": state.timestamp_min.isoformat(sep=" ") if state.timestamp_min else None,
            "local_max": state.timestamp_max.isoformat(sep=" ") if state.timestamp_max else None,
            "source_timezone": "unconfirmed",
        },
        "empty_post_content": state.empty_post_content,
        "possible_excel_truncation": state.possible_truncation,
        "duplicate_event_ids": state.duplicate_event_ids,
    }
    fingerprint = _dataset_fingerprint(dataset_version, updated_files, aggregate)
    tooling_after = _tooling_hashes()
    if tooling_before != tooling_after:
        raise ResearchStageError("TOOLING_CHANGED_DURING_AUDIT")
    manifest = {
        "manifest_version": MANIFEST_VERSION,
        "schema_version": SCHEMA_VERSION,
        "dataset_version": dataset_version,
        "dataset_fingerprint": fingerprint,
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "tooling": tooling_before,
        "source_files": [item.manifest_dict() for item in updated_files],
        "aggregate": aggregate,
        "data_semantics": {
            "snapshot_only": True,
            "operation": "snapshot",
            "export_time_utc": None,
            "available_from": None,
            "canonical_order": ["source_occurred_at_utc", "event_id"],
            "excluded_leakage_fields": [
                "comment_like_count",
                "unite_degree",
                "unite_degree_rank",
                "user_score",
                "user_level",
            ],
        },
        "privacy": {
            "classification": "pseudonymized_source_not_anonymous",
            "content_storage": "omitted_pending_local_clearance",
            "member_identity_fields_storage": "forbidden",
            "pii_candidate_counts_k_suppressed": _suppress_small_counts(
                state.pii_candidate_counts
            ),
            "external_api_calls": 0,
        },
        "projection_gate": {
            "learner_profile_projection_allowed": False,
            "reason": "SOURCE_KIND_UNCLASSIFIED_AND_SNAPSHOT_AVAILABILITY_UNKNOWN",
        },
    }
    return manifest, state


def _default_output_root() -> Path:
    local_app_data = os.environ.get("LOCALAPPDATA")
    if not local_app_data:
        raise ResearchStageError("LOCAL_OUTPUT_ROOT_REQUIRED")
    canonical_local = Path(local_app_data).expanduser().resolve()
    if _is_unc(canonical_local):
        raise ResearchStageError("UNC_OUTPUT_FORBIDDEN")
    configured = os.environ.get(OUTPUT_ROOT_ENV)
    if configured:
        output_root = Path(configured).expanduser().resolve()
        if _is_unc(output_root) or not _contains(canonical_local, output_root):
            raise ResearchStageError("OUTPUT_ROOT_OUTSIDE_LOCALAPPDATA")
        return output_root
    return (canonical_local / "cmooc-research" / "p1c").resolve()


def _validate_output_path(path_value: str, input_root: Path, *, must_not_exist: bool) -> Path:
    path = Path(path_value).expanduser()
    if _is_unc(path):
        raise ResearchStageError("UNC_OUTPUT_FORBIDDEN")
    output_root = _default_output_root()
    repo_root = _repo_root().resolve()
    resolved = path.resolve(strict=False)
    if _is_unc(output_root) or _is_unc(resolved):
        raise ResearchStageError("UNC_OUTPUT_FORBIDDEN")
    if not _contains(output_root, resolved):
        raise ResearchStageError("OUTPUT_OUTSIDE_LOCAL_ROOT")
    if _contains(repo_root, resolved) or _contains(input_root, resolved):
        raise ResearchStageError("SYNCED_OUTPUT_FORBIDDEN")
    protected = {
        (repo_root / "prisma" / "dev.db").resolve(),
        (repo_root / "prisma" / "test.db").resolve(),
    }
    if resolved in protected or resolved.name.lower() in {"dev.db", "test.db"}:
        raise ResearchStageError("BUSINESS_DATABASE_FORBIDDEN")
    if must_not_exist and resolved.exists():
        raise ResearchStageError("OUTPUT_ALREADY_EXISTS")
    resolved.parent.mkdir(parents=True, exist_ok=True)
    return resolved


def _write_immutable_json(path: Path, payload: Mapping[str, Any]) -> None:
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, indent=2) + "\n"
    try:
        with path.open("x", encoding="utf-8", newline="\n") as handle:
            handle.write(encoded)
    except FileExistsError as exc:
        raise ResearchStageError("OUTPUT_ALREADY_EXISTS") from exc


def _load_manifest(path: Path) -> dict[str, Any]:
    try:
        manifest = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ResearchStageError("MANIFEST_INVALID") from exc
    if (
        manifest.get("manifest_version") != MANIFEST_VERSION
        or manifest.get("schema_version") != SCHEMA_VERSION
        or not manifest.get("dataset_fingerprint")
    ):
        raise ResearchStageError("MANIFEST_INVALID")
    return manifest


def _load_hmac_key() -> tuple[bytes, str]:
    encoded = os.environ.get(HMAC_ENV, "")
    key_id = os.environ.get(HMAC_KEY_ID_ENV, "").strip()
    if not encoded or not key_id:
        raise ResearchStageError("HMAC_KEY_REQUIRED")
    try:
        key = base64.b64decode(encoded, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise ResearchStageError("HMAC_KEY_INVALID") from exc
    if len(key) < 32:
        raise ResearchStageError("HMAC_KEY_TOO_SHORT")
    if not re.fullmatch(r"[A-Za-z0-9._-]{1,64}", key_id):
        raise ResearchStageError("HMAC_KEY_ID_INVALID")
    return key, key_id


class Pseudonymizer:
    def __init__(self, key: bytes):
        self._key = key

    def token(self, domain: str, *parts: str, prefix: str) -> str:
        payload = "\x1f".join((domain, *parts)).encode("utf-8")
        digest = hmac.new(self._key, payload, hashlib.sha256).hexdigest()
        return f"{prefix}_{digest}"

    def course(self, course_id: str) -> str:
        return self.token("course", course_id, prefix="course")

    def actor(self, course_id: str, actor_id: str) -> str:
        return self.token("actor", course_id, actor_id, prefix="actor")

    def post(self, course_id: str, module_id: str) -> str:
        return self.token("post", course_id, module_id, prefix="post")

    def thread(self, course_id: str, module_id: str) -> str:
        return self.token("thread", course_id, module_id, prefix="thread")

    def reply(self, course_id: str, comment_id: str) -> str:
        return self.token("reply", course_id, comment_id, prefix="reply")

    def content(self, event_type: str, content: str) -> str:
        return self.token("content", event_type, content, prefix="content")

    def quarantine(self, logical_name: str, entity_type: str, source_id: str, reason: str) -> str:
        return self.token(
            "quarantine", logical_name, entity_type, source_id, reason, prefix="quarantine"
        )


def projection_decision(
    course_role: str,
    role_authority: str,
    source_kind: str,
    source_authority: str,
    timezone_confirmed: bool,
) -> tuple[bool, str]:
    if not timezone_confirmed:
        return False, "TIMEZONE_UNKNOWN"
    if course_role != "learner":
        return False, "ROLE_NOT_LEARNER"
    if role_authority != "platform_roster":
        return False, "ROLE_AUTHORITY_INSUFFICIENT"
    if source_kind != "human":
        return False, "SOURCE_NOT_CONFIRMED_HUMAN"
    if source_authority != "confirmed_dataset_profile":
        return False, "SOURCE_AUTHORITY_INSUFFICIENT"
    return True, "ELIGIBLE"


def _manifest_sources_match(expected: Mapping[str, Any], actual: Mapping[str, Any]) -> bool:
    keys = (
        "schema_version",
        "dataset_version",
        "dataset_fingerprint",
        "source_files",
        "tooling",
    )
    return all(expected.get(key) == actual.get(key) for key in keys)


def _insert_quarantine(
    connection: sqlite3.Connection,
    pseudonymizer: Pseudonymizer,
    dataset_fingerprint: str,
    logical_name: str,
    entity_type: str,
    source_id: str,
    reason: str,
    severity: str = "exclude",
) -> None:
    record_id = pseudonymizer.quarantine(logical_name, entity_type, source_id, reason)
    entity_ref = pseudonymizer.token(
        "quarantine_entity", logical_name, entity_type, source_id, prefix="entity"
    )
    connection.execute(
        """
        INSERT OR IGNORE INTO quarantine_record
          (record_id, dataset_fingerprint, source_file, entity_type, entity_ref, reason_code, severity)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (record_id, dataset_fingerprint, logical_name, entity_type, entity_ref, reason, severity),
    )


def _ensure_actor(
    connection: sqlite3.Connection,
    pseudonymizer: Pseudonymizer,
    course_raw: str,
    actor_raw: str,
    role: str,
    role_authority: str,
    source_kind: str,
    source_authority: str,
) -> tuple[str, str]:
    course_id = pseudonymizer.course(course_raw)
    actor_id = pseudonymizer.actor(course_raw, actor_raw)
    connection.execute(
        """
        INSERT OR IGNORE INTO research_actor
          (actor_id, course_id, display_name, course_role, role_authority, source_kind, source_authority)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (
            actor_id,
            course_id,
            f"Actor-{actor_id[-10:]}",
            role,
            role_authority,
            source_kind,
            source_authority,
        ),
    )
    return course_id, actor_id


def _insert_platform_role_claim(
    connection: sqlite3.Connection,
    pseudonymizer: Pseudonymizer,
    dataset_fingerprint: str,
    course_raw: str,
    actor_raw: str,
    role: str,
) -> None:
    course_id = pseudonymizer.course(course_raw)
    actor_id = pseudonymizer.actor(course_raw, actor_raw)
    claim_id = pseudonymizer.token(
        "role_claim",
        dataset_fingerprint,
        course_raw,
        actor_raw,
        "platform_role",
        role,
        prefix="claim",
    )
    connection.execute(
        """
        INSERT OR IGNORE INTO actor_role_claim
          (claim_id, dataset_fingerprint, course_id, actor_id, role_dimension,
           role_label, authority, confidence, temporal_scope, valid_from_utc,
           valid_until_utc, eligible_as_ground_truth)
        VALUES (?, ?, ?, ?, 'platform_role', ?, 'platform_roster', 1.0,
                'course_snapshot', NULL, NULL, 1)
        """,
        (claim_id, dataset_fingerprint, course_id, actor_id, role),
    )


def stage_dataset(
    *,
    input_root_value: str,
    manifest_path_value: str,
    database_path_value: str,
    source_timezone_name: str,
    source_profile: str,
    source_profile_confirmation: str | None,
    apply: bool,
) -> dict[str, Any]:
    if not apply:
        raise ResearchStageError("APPLY_CONFIRMATION_REQUIRED")
    try:
        source_timezone = ZoneInfo(source_timezone_name)
    except (ZoneInfoNotFoundError, ValueError) as exc:
        raise ResearchStageError("SOURCE_TIMEZONE_INVALID") from exc
    if source_profile not in {"unclassified", "historical-human-forum-export-v1"}:
        raise ResearchStageError("SOURCE_PROFILE_INVALID")
    if (
        source_profile == "historical-human-forum-export-v1"
        and source_profile_confirmation != "historical-human-forum-export-v1"
    ):
        raise ResearchStageError("SOURCE_PROFILE_CONFIRMATION_REQUIRED")

    input_root = _resolve_input_root(input_root_value)
    manifest_path = _validate_output_path(manifest_path_value, input_root, must_not_exist=False)
    manifest = _load_manifest(manifest_path)
    key, key_id = _load_hmac_key()
    database_path = _validate_output_path(database_path_value, input_root, must_not_exist=True)
    fresh_manifest, state = audit_dataset(input_root_value, str(manifest["dataset_version"]))
    if not _manifest_sources_match(manifest, fresh_manifest):
        raise ResearchStageError("MANIFEST_SOURCE_MISMATCH")

    pseudonymizer = Pseudonymizer(key)
    _, specs = _load_specs()
    source_paths = _resolve_source_files(input_root, specs)
    source_kind = "human" if source_profile == "historical-human-forum-export-v1" else "unknown"
    source_authority = (
        "confirmed_dataset_profile"
        if source_profile == "historical-human-forum-export-v1"
        else "unclassified"
    )
    fingerprint = str(manifest["dataset_fingerprint"])
    run_id = f"p1c_{fingerprint[:24]}"
    temp_path = database_path.with_name(f".{database_path.name}.{uuid.uuid4().hex}.tmp")
    connection: sqlite3.Connection | None = None
    try:
        connection = sqlite3.connect(temp_path)
        connection.execute("PRAGMA foreign_keys = ON")
        connection.execute("PRAGMA journal_mode = DELETE")
        connection.execute("PRAGMA synchronous = FULL")
        for migration_path in _migration_paths():
            connection.executescript(migration_path.read_text(encoding="utf-8"))
        connection.execute("BEGIN IMMEDIATE")
        connection.execute(
            """
            INSERT INTO dataset_manifest
              (dataset_fingerprint, dataset_version, schema_version, import_run_id,
               imported_at_utc, source_timezone, snapshot_only, export_time_utc,
               available_from_policy, content_policy, source_profile, key_id,
               importer_sha256, source_schema_sha256, migration_sha256,
               python_version, openpyxl_version, et_xmlfile_version, tzdata_version,
               sqlite_version, tooling_fingerprint,
               learner_projection_enabled)
            VALUES (?, ?, ?, ?, ?, ?, 1, NULL, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0)
            """,
            (
                fingerprint,
                manifest["dataset_version"],
                SCHEMA_VERSION,
                run_id,
                datetime.now(timezone.utc).isoformat(),
                source_timezone_name,
                "unknown_for_final_snapshot",
                "omitted_pending_local_clearance",
                source_profile,
                key_id,
                manifest["tooling"]["importer_sha256"],
                manifest["tooling"]["source_schema_sha256"],
                manifest["tooling"]["migration_sha256"],
                manifest["tooling"]["python_version"],
                manifest["tooling"]["openpyxl_version"],
                manifest["tooling"]["et_xmlfile_version"],
                manifest["tooling"]["tzdata_version"],
                manifest["tooling"]["sqlite_version"],
                manifest["tooling"]["tooling_fingerprint"],
            ),
        )
        for source_file in manifest["source_files"]:
            connection.execute(
                """
                INSERT INTO source_file
                  (dataset_fingerprint, logical_name, relative_path, sha256, byte_size,
                   worksheet, business_rows, accepted_rows, quarantined_rows)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    fingerprint,
                    source_file["logical_name"],
                    source_file["relative_path"],
                    source_file["sha256"],
                    source_file["byte_size"],
                    source_file["worksheet"],
                    source_file["business_rows"],
                    source_file["accepted_rows"],
                    source_file["quarantined_rows"],
                ),
            )

        for spec in specs:
            if spec.entity_type != "course":
                continue
            for row in _iter_rows(source_paths[spec.logical_name], spec):
                course_raw = _source_id(row.get("course_id"))
                if not course_raw:
                    _insert_quarantine(
                        connection,
                        pseudonymizer,
                        fingerprint,
                        spec.logical_name,
                        "course",
                        "missing",
                        "MISSING_COURSE_ID",
                    )
                    continue
                connection.execute(
                    "INSERT OR IGNORE INTO research_course (course_id, dataset_fingerprint, source_year) VALUES (?, ?, ?)",
                    (pseudonymizer.course(course_raw), fingerprint, spec.year),
                )

        for spec in specs:
            if spec.entity_type != "course_membership":
                continue
            for row in _iter_rows(source_paths[spec.logical_name], spec):
                course_raw = _source_id(row.get("course_id"))
                actor_raw = _source_id(row.get("user_id"))
                role_raw = _text(row.get("member_type")).strip().upper()
                if not course_raw or not actor_raw:
                    _insert_quarantine(
                        connection,
                        pseudonymizer,
                        fingerprint,
                        spec.logical_name,
                        "membership",
                        actor_raw or "missing",
                        "MISSING_ACTOR_ID",
                    )
                    continue
                role = ROLE_MAP.get(role_raw, "unknown")
                if role == "unknown":
                    _insert_quarantine(
                        connection,
                        pseudonymizer,
                        fingerprint,
                        spec.logical_name,
                        "membership",
                        actor_raw,
                        "UNKNOWN_COURSE_ROLE",
                    )
                course_id, actor_id = _ensure_actor(
                    connection,
                    pseudonymizer,
                    course_raw,
                    actor_raw,
                    role,
                    "platform_roster",
                    source_kind,
                    source_authority,
                )
                _insert_platform_role_claim(
                    connection,
                    pseudonymizer,
                    fingerprint,
                    course_raw,
                    actor_raw,
                    role,
                )
                connection.execute(
                    """
                    INSERT OR IGNORE INTO course_membership
                      (course_id, actor_id, course_role, role_authority, valid_from,
                       valid_until, membership_version)
                    VALUES (?, ?, ?, 'platform_roster', NULL, NULL, ?)
                    """,
                    (course_id, actor_id, role, fingerprint),
                )

        event_counts: Counter[str] = Counter()
        for spec in specs:
            if spec.entity_type != "post":
                continue
            for row in _iter_rows(source_paths[spec.logical_name], spec):
                course_raw = _source_id(row.get("course_id"))
                module_raw = _source_id(row.get("module_id"))
                actor_raw = _source_id(row.get("created_by"))
                if not course_raw or not module_raw or not actor_raw:
                    _insert_quarantine(
                        connection,
                        pseudonymizer,
                        fingerprint,
                        spec.logical_name,
                        "post",
                        module_raw or "missing",
                        "MISSING_REQUIRED_ID",
                    )
                    continue
                role = state.memberships.get((course_raw, actor_raw), "unknown")
                role_authority = "platform_roster" if role != "unknown" else "unknown"
                course_id, actor_id = _ensure_actor(
                    connection,
                    pseudonymizer,
                    course_raw,
                    actor_raw,
                    role,
                    role_authority,
                    source_kind,
                    source_authority,
                )
                event_id = pseudonymizer.post(course_raw, module_raw)
                occurred_at = _to_utc(row.get("created_time"), source_timezone)
                eligible, reason = projection_decision(
                    role, role_authority, source_kind, source_authority, True
                )
                if occurred_at is None:
                    eligible, reason = False, "INVALID_TIMESTAMP"
                elif eligible:
                    eligible, reason = False, "SNAPSHOT_AVAILABILITY_UNKNOWN"
                connection.execute(
                    """
                    INSERT INTO forum_event
                      (event_id, dataset_fingerprint, course_id, thread_id, post_id,
                       parent_event_id, actor_id, target_actor_id, event_type, operation,
                       course_role, role_authority, source_kind, source_authority,
                       source_occurred_at_utc, source_time_raw, source_timezone,
                       timestamp_precision, available_from_utc, projection_eligible,
                       projection_reason)
                    VALUES (?, ?, ?, ?, ?, NULL, ?, NULL, 'post', 'snapshot', ?, ?, ?, ?,
                            ?, ?, ?, 'second', NULL, ?, ?)
                    """,
                    (
                        event_id,
                        fingerprint,
                        course_id,
                        pseudonymizer.thread(course_raw, module_raw),
                        event_id,
                        actor_id,
                        role,
                        role_authority,
                        source_kind,
                        source_authority,
                        occurred_at,
                        _timestamp_raw(row.get("created_time")),
                        source_timezone_name,
                        int(eligible),
                        reason,
                    ),
                )
                title = _text(row.get("module_title"))
                body = _text(row.get("module_content"))
                content_payload = "\x1e".join((title, body))
                truncated = int(max(len(title), len(body)) >= MAX_CELL_LENGTH)
                normalized_char_count, paragraph_count, url_count = _content_structure(body)
                connection.execute(
                    """
                    INSERT INTO forum_event_content
                      (event_id, content_hmac, content_redacted, storage_policy,
                       content_empty, suspected_truncation, normalized_char_count,
                       paragraph_count, url_count, structure_signal_version)
                    VALUES (?, ?, NULL, 'omitted_pending_local_clearance', ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        event_id,
                        pseudonymizer.content("post", content_payload),
                        int(not content_payload.strip("\x1e ")),
                        truncated,
                        normalized_char_count,
                        paragraph_count,
                        url_count,
                        STRUCTURE_SIGNAL_VERSION,
                    ),
                )
                if truncated:
                    _insert_quarantine(
                        connection,
                        pseudonymizer,
                        fingerprint,
                        spec.logical_name,
                        "post",
                        module_raw,
                        "POSSIBLE_EXCEL_TRUNCATION",
                        "review",
                    )
                if occurred_at is None:
                    _insert_quarantine(
                        connection,
                        pseudonymizer,
                        fingerprint,
                        spec.logical_name,
                        "post",
                        module_raw,
                        "INVALID_TIMESTAMP",
                        "review",
                    )
                event_counts["posts"] += 1

        for spec in specs:
            if spec.entity_type != "reply":
                continue
            for row in _iter_rows(source_paths[spec.logical_name], spec):
                course_raw = _source_id(row.get("course_id"))
                comment_raw = _source_id(row.get("comment_id"))
                module_raw = _source_id(row.get("module_id"))
                actor_raw = _source_id(row.get("created_by"))
                source_ref = comment_raw or "missing"
                if not course_raw or not comment_raw or not module_raw or not actor_raw:
                    _insert_quarantine(
                        connection,
                        pseudonymizer,
                        fingerprint,
                        spec.logical_name,
                        "reply",
                        source_ref,
                        "MISSING_REQUIRED_ID",
                    )
                    continue
                key_tuple = (course_raw, comment_raw)
                if key_tuple not in state.canonical_comment_ids:
                    if (course_raw, module_raw) not in state.post_keys:
                        reason = "ORPHAN_POST"
                    else:
                        direct_parent_raw = _foreign_id(
                            row.get("comment_reply_child_id")
                        ) or _foreign_id(row.get("comment_reply_id"))
                        reason = (
                            "PARENT_NOT_CANONICAL"
                            if direct_parent_raw
                            and (course_raw, direct_parent_raw) in state.comment_ids
                            else "BROKEN_PARENT"
                        )
                    _insert_quarantine(
                        connection,
                        pseudonymizer,
                        fingerprint,
                        spec.logical_name,
                        "reply",
                        comment_raw,
                        reason,
                    )
                    continue
                role = state.memberships.get((course_raw, actor_raw), "unknown")
                role_authority = "platform_roster" if role != "unknown" else "unknown"
                course_id, actor_id = _ensure_actor(
                    connection,
                    pseudonymizer,
                    course_raw,
                    actor_raw,
                    role,
                    role_authority,
                    source_kind,
                    source_authority,
                )
                target_raw = _foreign_id(row.get("comment_reply_user_id"))
                target_id = None
                if target_raw:
                    _, target_id = _ensure_actor(
                        connection,
                        pseudonymizer,
                        course_raw,
                        target_raw,
                        state.memberships.get((course_raw, target_raw), "unknown"),
                        "platform_roster"
                        if (course_raw, target_raw) in state.memberships
                        else "unknown",
                        source_kind,
                        source_authority,
                    )
                is_reply = _is_reply(row.get("comment_reply"))
                direct_parent_raw = _foreign_id(row.get("comment_reply_child_id")) or _foreign_id(
                    row.get("comment_reply_id")
                )
                parent_id = (
                    pseudonymizer.reply(course_raw, direct_parent_raw)
                    if is_reply and direct_parent_raw
                    else pseudonymizer.post(course_raw, module_raw)
                )
                event_id = pseudonymizer.reply(course_raw, comment_raw)
                occurred_at = _to_utc(row.get("created_time"), source_timezone)
                eligible, reason = projection_decision(
                    role, role_authority, source_kind, source_authority, True
                )
                if occurred_at is None:
                    eligible, reason = False, "INVALID_TIMESTAMP"
                elif eligible:
                    eligible, reason = False, "SNAPSHOT_AVAILABILITY_UNKNOWN"
                connection.execute(
                    """
                    INSERT INTO forum_event
                      (event_id, dataset_fingerprint, course_id, thread_id, post_id,
                       parent_event_id, actor_id, target_actor_id, event_type, operation,
                       course_role, role_authority, source_kind, source_authority,
                       source_occurred_at_utc, source_time_raw, source_timezone,
                       timestamp_precision, available_from_utc, projection_eligible,
                       projection_reason)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'reply', 'snapshot', ?, ?, ?, ?,
                            ?, ?, ?, 'second', NULL, ?, ?)
                    """,
                    (
                        event_id,
                        fingerprint,
                        course_id,
                        pseudonymizer.thread(course_raw, module_raw),
                        pseudonymizer.post(course_raw, module_raw),
                        parent_id,
                        actor_id,
                        target_id,
                        role,
                        role_authority,
                        source_kind,
                        source_authority,
                        occurred_at,
                        _timestamp_raw(row.get("created_time")),
                        source_timezone_name,
                        int(eligible),
                        reason,
                    ),
                )
                content = _text(row.get("comment_content"))
                truncated = int(len(content) >= MAX_CELL_LENGTH)
                normalized_char_count, paragraph_count, url_count = _content_structure(content)
                connection.execute(
                    """
                    INSERT INTO forum_event_content
                      (event_id, content_hmac, content_redacted, storage_policy,
                       content_empty, suspected_truncation, normalized_char_count,
                       paragraph_count, url_count, structure_signal_version)
                    VALUES (?, ?, NULL, 'omitted_pending_local_clearance', ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        event_id,
                        pseudonymizer.content("reply", content),
                        int(not content.strip()),
                        truncated,
                        normalized_char_count,
                        paragraph_count,
                        url_count,
                        STRUCTURE_SIGNAL_VERSION,
                    ),
                )
                if truncated:
                    _insert_quarantine(
                        connection,
                        pseudonymizer,
                        fingerprint,
                        spec.logical_name,
                        "reply",
                        comment_raw,
                        "POSSIBLE_EXCEL_TRUNCATION",
                        "review",
                    )
                if occurred_at is None:
                    _insert_quarantine(
                        connection,
                        pseudonymizer,
                        fingerprint,
                        spec.logical_name,
                        "reply",
                        comment_raw,
                        "INVALID_TIMESTAMP",
                        "review",
                    )
                event_counts["comments"] += 1

        for source_file in manifest["source_files"]:
            current_hash = _sha256_file(source_paths[source_file["logical_name"]])
            if current_hash != source_file["sha256"]:
                raise ResearchStageError("SOURCE_CHANGED_DURING_STAGE")
        if manifest["tooling"] != _tooling_hashes():
            raise ResearchStageError("TOOLING_CHANGED_DURING_STAGE")

        staged_quarantine_counts = {
            reason: count
            for reason, count in connection.execute(
                "SELECT reason_code, COUNT(*) FROM quarantine_record GROUP BY reason_code ORDER BY reason_code"
            ).fetchall()
        }
        integrity = connection.execute("PRAGMA integrity_check").fetchone()[0]
        foreign_keys = connection.execute("PRAGMA foreign_key_check").fetchall()
        if integrity != "ok" or foreign_keys:
            raise ResearchStageError("STAGING_DATABASE_INTEGRITY_FAILED")
        connection.commit()
        connection.close()
        connection = None
        try:
            os.link(temp_path, database_path)
        except FileExistsError as exc:
            raise ResearchStageError("OUTPUT_ALREADY_EXISTS") from exc
        except OSError as exc:
            raise ResearchStageError("ATOMIC_PUBLISH_UNAVAILABLE") from exc
        try:
            temp_path.unlink()
        except OSError:
            # The final hard link is already complete; a stale temp name is harmless.
            pass
        return {
            "ok": True,
            "mode": "stage",
            "dataset_fingerprint": fingerprint,
            "source_timezone": source_timezone_name,
            "source_profile": source_profile,
            "projection_enabled": False,
            "event_counts": dict(sorted(event_counts.items())),
            "quarantine_counts": staged_quarantine_counts,
            "external_api_calls": 0,
            "content_storage": "omitted_pending_local_clearance",
        }
    except ResearchStageError:
        if connection is not None:
            connection.rollback()
            connection.close()
        if temp_path.exists():
            temp_path.unlink()
        raise
    except Exception as exc:
        if connection is not None:
            connection.rollback()
            connection.close()
        if temp_path.exists():
            temp_path.unlink()
        raise ResearchStageError("STAGING_FAILED") from exc


def _safe_summary(manifest: Mapping[str, Any], *, mode: str) -> dict[str, Any]:
    return {
        "ok": True,
        "mode": mode,
        "dataset_version": manifest["dataset_version"],
        "dataset_fingerprint": manifest["dataset_fingerprint"],
        "aggregate": manifest["aggregate"],
        "projection_gate": manifest["projection_gate"],
        "privacy": manifest["privacy"],
    }


def _build_parser() -> argparse.ArgumentParser:
    class SafeArgumentParser(argparse.ArgumentParser):
        def error(self, message: str) -> None:
            raise ResearchStageError("CLI_ARGUMENT_INVALID")

    parser = SafeArgumentParser(description="Local-only cMOOC research staging")
    subparsers = parser.add_subparsers(dest="command", required=True)

    audit = subparsers.add_parser("audit", help="read-only source audit")
    audit.add_argument("--input-root", required=True)
    audit.add_argument("--dataset-version", required=True)
    audit.add_argument("--manifest-out")

    stage = subparsers.add_parser("stage", help="create a new isolated research database")
    stage.add_argument("--input-root", required=True)
    stage.add_argument("--manifest", required=True)
    stage.add_argument("--db", required=True)
    stage.add_argument("--timezone", required=True)
    stage.add_argument(
        "--source-profile",
        choices=("unclassified", "historical-human-forum-export-v1"),
        default="unclassified",
    )
    stage.add_argument("--confirm-source-profile")
    stage.add_argument("--apply", action="store_true")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    try:
        _deny_network()
        args = _build_parser().parse_args(argv)
        if args.command == "audit":
            manifest, _ = audit_dataset(args.input_root, args.dataset_version)
            if args.manifest_out:
                input_root = _resolve_input_root(args.input_root)
                output_path = _validate_output_path(
                    args.manifest_out, input_root, must_not_exist=True
                )
                _write_immutable_json(output_path, manifest)
            result = _safe_summary(manifest, mode="audit")
        else:
            result = stage_dataset(
                input_root_value=args.input_root,
                manifest_path_value=args.manifest,
                database_path_value=args.db,
                source_timezone_name=args.timezone,
                source_profile=args.source_profile,
                source_profile_confirmation=args.confirm_source_profile,
                apply=args.apply,
            )
        print(json.dumps(result, ensure_ascii=False, sort_keys=True))
        return 0
    except ResearchStageError as exc:
        print(json.dumps({"ok": False, "error_code": exc.code}, sort_keys=True), file=sys.stderr)
        return 2
    except Exception:
        print(json.dumps({"ok": False, "error_code": "INTERNAL_ERROR"}, sort_keys=True), file=sys.stderr)
        return 3


if __name__ == "__main__":
    raise SystemExit(main())
