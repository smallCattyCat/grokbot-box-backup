#!/usr/bin/env python3
"""Load a frozen P3A decision-reference batch into memory at its as-of boundary."""

from __future__ import annotations

import base64
import binascii
import csv
import hashlib
import hmac
import io
import json
import os
import re
import tempfile
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook

from decision_reference_integrity import (
    ReferenceIntegrityError,
    reference_integrity,
    verify_reference_integrity,
)
from research_review_content import (
    MAX_CELL_LENGTH,
    ReviewContentError,
    _identity_lexicon,
    _plain_text,
    _redact,
)


SELECTION_SCHEMA_VERSION = "cmooc-decision-reference-selection-v1"
LOCK_SCHEMA_VERSION = "cmooc-decision-reference-lock-v1"
CONTENT_REVISION_STATUS = "final_export_revision_unknown"
MAX_REFERENCE_ITEMS = 60
STABLE_KEY_PATTERN = re.compile(r"[0-9a-f]{24}")
REFERENCE_ID_PATTERN = re.compile(r"DR[0-9]{3}")
SELECTION_LOCK_CLAIM_FILENAME = ".selection-lock.create"
SELECTION_FIELDS = (
    "reference_id",
    "snapshot_key",
    "thread_key",
    "course_key",
    "author_key",
    "year",
    "decision_at",
    "sampling_stratum",
    "sampling_reason",
    "intervene",
    "main_function",
    "target_object",
)


class DecisionReviewContentError(RuntimeError):
    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


@dataclass(frozen=True)
class SelectionItem:
    reference_id: str
    snapshot_key: str
    thread_key: str
    course_key: str
    author_key: str
    year: int
    decision_at: str
    display_order: int


@dataclass(frozen=True)
class SelectionContract:
    schema_version: str
    selection_fingerprint: str
    item_count: int
    items: tuple[SelectionItem, ...]


@dataclass(frozen=True)
class DecisionReviewEvent:
    event_at: str
    event_kind: str
    actor_label: str
    body: str
    reply_to_actor_label: str | None


@dataclass(frozen=True)
class DecisionReviewContent:
    reference_id: str
    snapshot_key: str
    thread_key: str
    display_order: int
    decision_at: str
    thread_created_at: str
    title: str
    body: str
    events: tuple[DecisionReviewEvent, ...]
    suspected_truncation: bool
    redaction_count: int
    content_revision_status: str


@dataclass(frozen=True)
class _RawSelectedPost:
    selection: SelectionItem
    course_id: str
    module_id: str
    author_id: str
    created_at: datetime
    raw_title: str
    raw_body: str


@dataclass(frozen=True)
class _RawComment:
    comment_id: str
    author_id: str
    created_at: datetime
    raw_body: str
    is_reply: bool
    parent_comment_id: str


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[1]


def research_home() -> Path:
    configured = os.environ.get("CMOOC_RESEARCH_HOME")
    if configured:
        return Path(configured).expanduser().resolve()
    local_app_data = os.environ.get("LOCALAPPDATA")
    if not local_app_data:
        raise DecisionReviewContentError("LOCAL_RESEARCH_HOME_REQUIRED")
    return Path(local_app_data).resolve() / "cmooc-research"


def default_input_root() -> Path:
    configured = os.environ.get("CMOOC_RESEARCH_INPUT_ROOT")
    if configured:
        return Path(configured).expanduser().resolve()
    return _repo_root().parent.parent / "cmooc研究数据24-25"


def default_selection_file() -> Path:
    return (
        research_home()
        / "p3a"
        / "decision-snapshot-v1"
        / "reference"
        / "decision-reference-60.csv"
    )


def default_key_path() -> Path:
    return research_home() / "secrets" / "decision-snapshot-v1.key"


def default_output_root() -> Path:
    return research_home() / "p3b" / "decision-reference-v1"


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _normalize_id(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_timestamp(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    text = _normalize_id(value)
    if not text:
        return None
    try:
        return datetime.fromisoformat(
            text.replace("/", "-").replace("T", " ").strip()
        ).replace(tzinfo=None)
    except ValueError:
        return None


def _iso(value: datetime) -> str:
    return value.isoformat(timespec="seconds")


def _stable_key(key: bytes, namespace: str, *parts: object) -> str:
    message = "\x1f".join([namespace, *(str(part) for part in parts)])
    return hmac.new(
        key, message.encode("utf-8"), hashlib.sha256
    ).hexdigest()[:24]


def load_decision_key(path: Path) -> bytes:
    if not path.is_file():
        raise DecisionReviewContentError("DECISION_KEY_REQUIRED")
    try:
        key = base64.b64decode(
            path.read_text(encoding="ascii").strip(), validate=True
        )
    except (OSError, UnicodeError, binascii.Error, ValueError) as exc:
        raise DecisionReviewContentError("DECISION_KEY_INVALID") from exc
    if len(key) != 32:
        raise DecisionReviewContentError("DECISION_KEY_INVALID")
    return key


def load_selection_contract(path: Path) -> SelectionContract:
    path = path.expanduser().resolve()
    if not path.is_file():
        raise DecisionReviewContentError("SELECTION_FILE_REQUIRED")
    try:
        selection_bytes = path.read_bytes()
        fingerprint = hashlib.sha256(selection_bytes).hexdigest()
        selection_text = selection_bytes.decode("utf-8-sig")
        reader = csv.DictReader(
            io.StringIO(selection_text, newline="")
        )
        if tuple(reader.fieldnames or ()) != SELECTION_FIELDS:
            raise DecisionReviewContentError("SELECTION_SCHEMA_INVALID")
        rows = list(reader)
    except (OSError, UnicodeError, csv.Error) as exc:
        raise DecisionReviewContentError("SELECTION_FILE_INVALID") from exc
    if not rows:
        raise DecisionReviewContentError("SELECTION_EMPTY")
    if len(rows) > MAX_REFERENCE_ITEMS:
        raise DecisionReviewContentError("SELECTION_LIMIT_EXCEEDED")

    seen_references: set[str] = set()
    seen_snapshots: set[str] = set()
    seen_threads: set[str] = set()
    items: list[SelectionItem] = []
    for display_order, row in enumerate(rows, start=1):
        reference_id = str(row.get("reference_id", ""))
        snapshot_key = str(row.get("snapshot_key", ""))
        thread_key = str(row.get("thread_key", ""))
        course_key = str(row.get("course_key", ""))
        author_key = str(row.get("author_key", ""))
        if (
            not REFERENCE_ID_PATTERN.fullmatch(reference_id)
            or any(
                not STABLE_KEY_PATTERN.fullmatch(value)
                for value in (
                    snapshot_key,
                    thread_key,
                    course_key,
                    author_key,
                )
            )
        ):
            raise DecisionReviewContentError("SELECTION_STABLE_KEY_INVALID")
        if (
            reference_id in seen_references
            or snapshot_key in seen_snapshots
            or thread_key in seen_threads
        ):
            raise DecisionReviewContentError("SELECTION_DUPLICATE_ITEM")
        if any(
            str(row.get(field, "")).strip()
            for field in ("intervene", "main_function", "target_object")
        ):
            raise DecisionReviewContentError(
                "SELECTION_LABELS_MUST_BE_EMPTY"
            )
        try:
            year = int(str(row.get("year", "")))
            decision_at_value = str(row.get("decision_at", ""))
            decision_at = datetime.fromisoformat(decision_at_value)
        except (TypeError, ValueError) as exc:
            raise DecisionReviewContentError(
                "SELECTION_DECISION_TIME_INVALID"
            ) from exc
        if (
            decision_at.tzinfo is not None
            or year not in {2024, 2025}
            or decision_at.year not in {year, year + 1}
            or _iso(decision_at) != decision_at_value
        ):
            raise DecisionReviewContentError(
                "SELECTION_DECISION_TIME_INVALID"
            )
        seen_references.add(reference_id)
        seen_snapshots.add(snapshot_key)
        seen_threads.add(thread_key)
        items.append(
            SelectionItem(
                reference_id=reference_id,
                snapshot_key=snapshot_key,
                thread_key=thread_key,
                course_key=course_key,
                author_key=author_key,
                year=year,
                decision_at=decision_at_value,
                display_order=display_order,
            )
        )
    return SelectionContract(
        schema_version=SELECTION_SCHEMA_VERSION,
        selection_fingerprint=fingerprint,
        item_count=len(items),
        items=tuple(items),
    )


def selection_contract_integrity(
    contract: SelectionContract,
) -> dict[str, str]:
    return reference_integrity(contract.items)


def verify_selection_manifest(
    selection_path: Path, contract: SelectionContract
) -> dict[str, Any]:
    manifest_path = selection_path.parent / "reference-manifest.json"
    if not manifest_path.is_file():
        raise DecisionReviewContentError(
            "REFERENCE_MANIFEST_REQUIRED"
        )
    try:
        manifest = json.loads(
            manifest_path.read_text(encoding="utf-8")
        )
    except (
        OSError,
        UnicodeError,
        json.JSONDecodeError,
    ) as exc:
        raise DecisionReviewContentError(
            "REFERENCE_MANIFEST_INVALID"
        ) from exc
    if (
        not isinstance(manifest, dict)
        or manifest.get("schema_version")
        != SELECTION_SCHEMA_VERSION
        or manifest.get("item_count") != contract.item_count
        or manifest.get("maximum_item_count")
        != MAX_REFERENCE_ITEMS
        or manifest.get("selection_frozen") is not True
    ):
        raise DecisionReviewContentError(
            "REFERENCE_MANIFEST_INVALID"
        )
    try:
        verify_reference_integrity(manifest, contract.items)
    except ReferenceIntegrityError as exc:
        raise DecisionReviewContentError(exc.code) from exc
    return manifest


def _source_paths(
    *,
    input_root: Path,
    selection_path: Path,
    source_audit_path: Path | None,
) -> dict[str, Path]:
    raw_root = input_root / "01_原始数据" / "数据统计"
    audit_path = (
        source_audit_path
        if source_audit_path is not None
        else selection_path.parent.parent / "source_audit.json"
    )
    try:
        audit = json.loads(audit_path.read_text(encoding="utf-8"))
        source_files = audit["source_files"]
    except (OSError, UnicodeError, json.JSONDecodeError, KeyError) as exc:
        raise DecisionReviewContentError("SOURCE_AUDIT_INVALID") from exc
    if not isinstance(source_files, dict):
        raise DecisionReviewContentError("SOURCE_AUDIT_INVALID")

    required_names = {
        f"{year}{suffix}"
        for year in (2024, 2025)
        for suffix in ("图文.xlsx", "图文评论.xlsx", "课程成员.xlsx")
    }
    if not required_names.issubset(source_files):
        raise DecisionReviewContentError("SOURCE_AUDIT_INCOMPLETE")
    resolved: dict[str, Path] = {}
    for name, metadata in source_files.items():
        if (
            not isinstance(name, str)
            or Path(name).name != name
            or not isinstance(metadata, dict)
            or not isinstance(metadata.get("sha256"), str)
        ):
            raise DecisionReviewContentError("SOURCE_AUDIT_INVALID")
        path = (raw_root / name).resolve()
        if (
            not path.is_file()
            or _file_sha256(path) != metadata["sha256"]
        ):
            raise DecisionReviewContentError("SOURCE_FILE_CHANGED")
        resolved[name] = path
    return resolved


def _header_indices(
    worksheet: Any, required: set[str]
) -> dict[str, int]:
    try:
        header = next(
            worksheet.iter_rows(min_row=1, max_row=1, values_only=True)
        )
    except StopIteration as exc:
        raise DecisionReviewContentError("SOURCE_COLUMNS_MISSING") from exc
    indices = {
        str(value).strip(): index for index, value in enumerate(header)
    }
    if not required.issubset(indices):
        raise DecisionReviewContentError("SOURCE_COLUMNS_MISSING")
    return indices


def _cell_value(cells: Sequence[Any], index: int) -> Any:
    return cells[index].value if index < len(cells) else None


def _formula_free(cells: Sequence[Any]) -> None:
    if any(cell.data_type == "f" for cell in cells):
        raise DecisionReviewContentError("XLSX_FORMULA_FORBIDDEN")


def _selected_posts(
    paths: Mapping[str, Path],
    contract: SelectionContract,
    key: bytes,
) -> dict[str, _RawSelectedPost]:
    selected_by_thread = {
        item.thread_key: item for item in contract.items
    }
    result: dict[str, _RawSelectedPost] = {}
    required = {
        "module_id",
        "course_id",
        "module_title",
        "module_content",
        "created_by",
        "created_time",
    }
    for year in (2024, 2025):
        path = paths[f"{year}图文.xlsx"]
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            worksheet = workbook[workbook.sheetnames[0]]
            indices = _header_indices(worksheet, required)
            for cells in worksheet.iter_rows(min_row=3, values_only=False):
                _formula_free(cells)
                course_id = _normalize_id(
                    _cell_value(cells, indices["course_id"])
                )
                module_id = _normalize_id(
                    _cell_value(cells, indices["module_id"])
                )
                if not course_id or not module_id:
                    continue
                thread_key = _stable_key(
                    key, "thread", year, course_id, module_id
                )
                item = selected_by_thread.get(thread_key)
                if item is None:
                    continue
                if item.year != year or item.reference_id in result:
                    raise DecisionReviewContentError(
                        "SELECTED_CONTENT_DUPLICATE"
                    )
                author_id = _normalize_id(
                    _cell_value(cells, indices["created_by"])
                )
                created_at = _parse_timestamp(
                    _cell_value(cells, indices["created_time"])
                )
                decision_at = datetime.fromisoformat(item.decision_at)
                if (
                    not author_id
                    or created_at is None
                    or created_at > decision_at
                    or _stable_key(
                        key,
                        "snapshot",
                        year,
                        course_id,
                        module_id,
                        item.decision_at,
                    )
                    != item.snapshot_key
                    or _stable_key(key, "course", year, course_id)
                    != item.course_key
                    or _stable_key(key, "author", author_id)
                    != item.author_key
                ):
                    raise DecisionReviewContentError(
                        "SELECTION_SOURCE_MISMATCH"
                    )
                raw_title = str(
                    _cell_value(cells, indices["module_title"]) or ""
                )
                raw_body = str(
                    _cell_value(cells, indices["module_content"]) or ""
                )
                result[item.reference_id] = _RawSelectedPost(
                    selection=item,
                    course_id=course_id,
                    module_id=module_id,
                    author_id=author_id,
                    created_at=created_at,
                    raw_title=raw_title,
                    raw_body=raw_body,
                )
        finally:
            workbook.close()
    if len(result) != contract.item_count:
        raise DecisionReviewContentError("SELECTED_CONTENT_NOT_FOUND")
    return result


def _asof_comments(
    paths: Mapping[str, Path],
    posts: Mapping[str, _RawSelectedPost],
) -> dict[str, list[_RawComment]]:
    raw_post_index = {
        (post.selection.year, post.course_id, post.module_id): reference_id
        for reference_id, post in posts.items()
    }
    result: dict[str, list[_RawComment]] = {
        reference_id: [] for reference_id in posts
    }
    seen: set[tuple[int, str, str, str]] = set()
    required = {
        "course_id",
        "comment_id",
        "module_id",
        "comment_content",
        "comment_reply",
        "comment_reply_id",
        "comment_reply_child_id",
        "created_by",
        "created_time",
    }
    for year in (2024, 2025):
        path = paths[f"{year}图文评论.xlsx"]
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            worksheet = workbook[workbook.sheetnames[0]]
            indices = _header_indices(worksheet, required)
            for cells in worksheet.iter_rows(min_row=3, values_only=False):
                _formula_free(cells)
                course_id = _normalize_id(
                    _cell_value(cells, indices["course_id"])
                )
                module_id = _normalize_id(
                    _cell_value(cells, indices["module_id"])
                )
                reference_id = raw_post_index.get(
                    (year, course_id, module_id)
                )
                if reference_id is None:
                    continue
                comment_id = _normalize_id(
                    _cell_value(cells, indices["comment_id"])
                )
                author_id = _normalize_id(
                    _cell_value(cells, indices["created_by"])
                )
                created_at = _parse_timestamp(
                    _cell_value(cells, indices["created_time"])
                )
                raw_key = (year, course_id, module_id, comment_id)
                if not comment_id or raw_key in seen:
                    raise DecisionReviewContentError(
                        "COMMENT_STABLE_KEY_INVALID"
                    )
                seen.add(raw_key)
                if not author_id or created_at is None:
                    raise DecisionReviewContentError(
                        "COMMENT_SOURCE_INVALID"
                    )
                post = posts[reference_id]
                decision_at = datetime.fromisoformat(
                    post.selection.decision_at
                )
                if (
                    created_at < post.created_at
                    or created_at > decision_at
                ):
                    continue
                parent = _normalize_id(
                    _cell_value(
                        cells, indices["comment_reply_child_id"]
                    )
                )
                if parent in {"", "-1"}:
                    parent = _normalize_id(
                        _cell_value(cells, indices["comment_reply_id"])
                    )
                if parent == "-1":
                    parent = ""
                result[reference_id].append(
                    _RawComment(
                        comment_id=comment_id,
                        author_id=author_id,
                        created_at=created_at,
                        raw_body=str(
                            _cell_value(
                                cells, indices["comment_content"]
                            )
                            or ""
                        ),
                        is_reply=(
                            _normalize_id(
                                _cell_value(
                                    cells, indices["comment_reply"]
                                )
                            )
                            == "1"
                            or bool(parent)
                        ),
                        parent_comment_id=parent,
                    )
                )
        finally:
            workbook.close()
    for comments in result.values():
        comments.sort(key=lambda item: (item.created_at, item.comment_id))
    return result


def _participant_label(index: int) -> str:
    if index < 26:
        return f"参与者 {chr(ord('A') + index)}"
    return f"参与者 {index + 1}"


def _render_content(
    *,
    post: _RawSelectedPost,
    comments: Sequence[_RawComment],
    identities: Sequence[str],
) -> DecisionReviewContent:
    title, title_redactions = _redact(
        _plain_text(post.raw_title), identities
    )
    body, body_redactions = _redact(
        _plain_text(post.raw_body), identities
    )
    actor_labels = {post.author_id: "主题作者"}
    for comment in comments:
        if comment.author_id not in actor_labels:
            actor_labels[comment.author_id] = _participant_label(
                len(actor_labels) - 1
            )
    author_by_comment = {
        comment.comment_id: comment.author_id for comment in comments
    }
    redaction_count = title_redactions + body_redactions
    suspected_truncation = max(
        len(post.raw_title), len(post.raw_body)
    ) >= MAX_CELL_LENGTH
    events: list[DecisionReviewEvent] = []
    decision_at = datetime.fromisoformat(post.selection.decision_at)
    for comment in comments:
        clean_body, count = _redact(
            _plain_text(comment.raw_body), identities
        )
        redaction_count += count
        suspected_truncation = (
            suspected_truncation
            or len(comment.raw_body) >= MAX_CELL_LENGTH
        )
        parent_author = author_by_comment.get(comment.parent_comment_id)
        event = DecisionReviewEvent(
            event_at=_iso(comment.created_at),
            event_kind="reply" if comment.is_reply else "comment",
            actor_label=actor_labels[comment.author_id],
            body=clean_body or "（内容为空）",
            reply_to_actor_label=(
                actor_labels.get(parent_author)
                if parent_author is not None
                else None
            ),
        )
        if datetime.fromisoformat(event.event_at) > decision_at:
            raise DecisionReviewContentError("ASOF_BOUNDARY_VIOLATION")
        events.append(event)
    return DecisionReviewContent(
        reference_id=post.selection.reference_id,
        snapshot_key=post.selection.snapshot_key,
        thread_key=post.selection.thread_key,
        display_order=post.selection.display_order,
        decision_at=post.selection.decision_at,
        thread_created_at=_iso(post.created_at),
        title=title or "（无标题）",
        body=body or "（正文为空）",
        events=tuple(events),
        suspected_truncation=suspected_truncation,
        redaction_count=redaction_count,
        content_revision_status=CONTENT_REVISION_STATUS,
    )


def load_decision_review_contents(
    *,
    selection_file_value: str,
    input_root_value: str,
    key_path_value: str,
    source_audit_value: str | None = None,
) -> tuple[SelectionContract, dict[str, DecisionReviewContent]]:
    selection_path = Path(selection_file_value).expanduser().resolve()
    input_root = Path(input_root_value).expanduser().resolve()
    key_path = Path(key_path_value).expanduser().resolve()
    source_audit_path = (
        Path(source_audit_value).expanduser().resolve()
        if source_audit_value
        else None
    )
    contract = load_selection_contract(selection_path)
    verify_selection_manifest(selection_path, contract)
    key = load_decision_key(key_path)
    paths = _source_paths(
        input_root=input_root,
        selection_path=selection_path,
        source_audit_path=source_audit_path,
    )
    try:
        identities = _identity_lexicon(
            [
                paths["2024课程成员.xlsx"],
                paths["2025课程成员.xlsx"],
            ]
        )
    except ReviewContentError as exc:
        raise DecisionReviewContentError(exc.code) from exc
    posts = _selected_posts(paths, contract, key)
    del key
    comments = _asof_comments(paths, posts)
    contents = {
        reference_id: _render_content(
            post=post,
            comments=comments[reference_id],
            identities=identities,
        )
        for reference_id, post in posts.items()
    }
    return contract, contents


def public_item_payload(
    content: DecisionReviewContent,
) -> dict[str, Any]:
    return {
        "reference_id": content.reference_id,
        "display_order": content.display_order,
        "decision_at": content.decision_at,
        "thread_created_at": content.thread_created_at,
        "title": content.title,
        "body": content.body,
        "events": [asdict(event) for event in content.events],
        "suspected_truncation": content.suspected_truncation,
        "redaction_count": content.redaction_count,
        "content_revision_status": content.content_revision_status,
    }


def _atomic_write_json(path: Path, payload: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(
        prefix=".decision-reference-lock-",
        suffix=".tmp",
        dir=path.parent,
    )
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as handle:
            json.dump(
                payload,
                handle,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            )
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp_name, path)
    finally:
        try:
            Path(temp_name).unlink(missing_ok=True)
        except OSError:
            pass


def ensure_selection_lock(
    output_root: Path, contract: SelectionContract
) -> Path:
    output_root = output_root.expanduser().resolve()
    output_root.mkdir(parents=True, exist_ok=True)
    lock_path = output_root / "selection-lock.json"
    claim_path = output_root / SELECTION_LOCK_CLAIM_FILENAME
    payload = {
        "schema_version": LOCK_SCHEMA_VERSION,
        "selection_fingerprint": contract.selection_fingerprint,
        "item_count": contract.item_count,
        "maximum_item_count": MAX_REFERENCE_ITEMS,
        **selection_contract_integrity(contract),
        "items": [
            {
                "reference_id": item.reference_id,
                "snapshot_key": item.snapshot_key,
                "thread_key": item.thread_key,
                "display_order": item.display_order,
            }
            for item in contract.items
        ],
    }

    def validate_existing() -> bool:
        if not lock_path.is_file():
            return False
        try:
            existing = json.loads(lock_path.read_text(encoding="utf-8"))
        except (OSError, UnicodeError, json.JSONDecodeError) as exc:
            raise DecisionReviewContentError(
                "SELECTION_LOCK_INVALID"
            ) from exc
        if existing != payload:
            raise DecisionReviewContentError("SELECTION_LOCK_MISMATCH")
        return True

    if validate_existing():
        return lock_path

    try:
        claim_descriptor = os.open(
            claim_path,
            os.O_CREAT | os.O_EXCL | os.O_WRONLY,
            0o600,
        )
    except FileExistsError as exc:
        raise DecisionReviewContentError(
            "SELECTION_LOCK_CREATION_IN_PROGRESS"
        ) from exc
    try:
        if validate_existing():
            return lock_path
        _atomic_write_json(lock_path, payload)
        return lock_path
    finally:
        os.close(claim_descriptor)
        try:
            claim_path.unlink(missing_ok=True)
        except OSError:
            pass
